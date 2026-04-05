from __future__ import annotations

import asyncio
import csv
import json
import os
import shutil
import sys
import tkinter as tk
import traceback
from collections import Counter
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from tkinter import filedialog

import flet as ft
import qrcode
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from check_excel_completeness import run_checks

if getattr(sys, "frozen", False):
    RESOURCE_DIR = Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent))
    APP_DIR = Path(sys.executable).resolve().parent
else:
    RESOURCE_DIR = Path(__file__).resolve().parent
    APP_DIR = RESOURCE_DIR


DATA_DIR = APP_DIR / "data"
DATA_CONFIG_DIR = DATA_DIR / "config"
DATA_DATABASE_DIR = DATA_DIR / "database"
DATA_DOCS_DIR = DATA_DIR / "docs"
DATA_LOG_DIR = DATA_DIR / "log"

LEGACY_LOG_FILE = APP_DIR / "ips-checker-error.log"
LOG_FILE = DATA_LOG_DIR / "ips-checker-error.log"
ICON_FILE = RESOURCE_DIR / "assets" / "app_icon.ico"
IPS_HTML_FILE = RESOURCE_DIR / "assets" / "ips.html"
LEGACY_CONFIG_FILE = APP_DIR / "ips-checker-config.json"
LEGACY_FOLLOW_UP_DB_FILE = APP_DIR / "ips-follow-up-database.csv"
CONFIG_FILE = DATA_CONFIG_DIR / "ips-checker-config.json"
RUNTIME_FIELD_FILE = DATA_CONFIG_DIR / "field_cell.txt"
FOLLOW_UP_DB_FILE = DATA_DATABASE_DIR / "ips-follow-up-database.csv"
CHECKER_GUIDE_FILE = DATA_DOCS_DIR / "ips-checker-user-guide.pdf"
GENERATOR_GUIDE_FILE = DATA_DOCS_DIR / "ips-generator-user-guide.pdf"
DEFAULT_WINDOW_WIDTH = 1180
DEFAULT_WINDOW_HEIGHT = 820


def write_error_log(context: str, exc: BaseException) -> Path:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log_entry = "\n".join(
        [
            f"[{timestamp}] {context}",
            f"Type: {type(exc).__name__}",
            f"Message: {exc}",
            "Traceback:",
            "".join(
                traceback.format_exception(type(exc), exc, exc.__traceback__)
            ).rstrip(),
            "-" * 80,
            "",
        ]
    )
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    with LOG_FILE.open("a", encoding="utf-8") as log_file:
        log_file.write(log_entry)
    return LOG_FILE


def migrate_file_if_missing(source_path: Path, destination_path: Path) -> None:
    if not source_path.exists() or destination_path.exists():
        return

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.move(str(source_path), str(destination_path))
    except Exception:
        shutil.copy2(source_path, destination_path)


def copy_file_if_newer(source_path: Path, destination_path: Path) -> None:
    if not source_path.exists():
        return

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    if destination_path.exists():
        try:
            if source_path.stat().st_mtime <= destination_path.stat().st_mtime:
                return
        except OSError:
            pass

    shutil.copy2(source_path, destination_path)


def get_default_field_file_path() -> Path:
    if RUNTIME_FIELD_FILE.exists():
        return RUNTIME_FIELD_FILE

    bundled_field_file = RESOURCE_DIR / "field_cell.txt"
    if bundled_field_file.exists():
        return bundled_field_file

    return APP_DIR / "field_cell.txt"


def initialize_app_storage() -> None:
    for directory in (
        DATA_DIR,
        DATA_CONFIG_DIR,
        DATA_DATABASE_DIR,
        DATA_DOCS_DIR,
        DATA_LOG_DIR,
    ):
        directory.mkdir(parents=True, exist_ok=True)

    migrate_file_if_missing(LEGACY_LOG_FILE, LOG_FILE)
    migrate_file_if_missing(LEGACY_CONFIG_FILE, CONFIG_FILE)
    migrate_file_if_missing(LEGACY_FOLLOW_UP_DB_FILE, FOLLOW_UP_DB_FILE)
    copy_file_if_newer(RESOURCE_DIR / "field_cell.txt", RUNTIME_FIELD_FILE)

    source_docs_dir = RESOURCE_DIR / "docs"
    copy_file_if_newer(
        source_docs_dir / "ips-checker-user-guide.pdf", CHECKER_GUIDE_FILE
    )
    copy_file_if_newer(
        source_docs_dir / "ips-generator-user-guide.pdf", GENERATOR_GUIDE_FILE
    )

    normalize_migrated_config_defaults()


def normalize_migrated_config_defaults() -> None:
    if not CONFIG_FILE.exists():
        return

    try:
        config_payload = json.loads(CONFIG_FILE.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        write_error_log("Gagal membaca config untuk normalisasi migrasi", exc)
        return

    if not isinstance(config_payload, dict):
        return

    normalized_payload = dict(config_payload)
    normalized_payload["follow_up_db_path"] = str(FOLLOW_UP_DB_FILE)

    current_field_file = str(normalized_payload.get("field_file") or "").strip()
    default_field_file = get_default_field_file_path()
    resolved_current_field_file = (
        Path(current_field_file) if current_field_file else None
    )
    should_reset_field_file = not current_field_file
    if (
        resolved_current_field_file is not None
        and not resolved_current_field_file.exists()
    ):
        should_reset_field_file = True

    if should_reset_field_file:
        normalized_payload["field_file"] = str(default_field_file)

    if normalized_payload == config_payload:
        return

    try:
        CONFIG_FILE.write_text(
            json.dumps(normalized_payload, indent=2), encoding="utf-8"
        )
    except Exception as exc:
        write_error_log("Gagal menyimpan config hasil normalisasi migrasi", exc)


def sanitize_window_size(value: object, fallback: int) -> int:
    try:
        parsed_value = int(float(value))
    except TypeError, ValueError:
        return fallback
    return parsed_value if parsed_value >= 600 else fallback


def load_app_config(
    default_target_path: Path, default_field_file: Path
) -> dict[str, object]:
    config = {
        "target_path": str(default_target_path),
        "field_file": str(default_field_file),
        "sheet_mode": "auto",
        "sheet_name": None,
        "csv_output": str(APP_DIR / "report.csv"),
        "follow_up_db_path": str(FOLLOW_UP_DB_FILE),
        "result_view": "check-results",
        "window_width": DEFAULT_WINDOW_WIDTH,
        "window_height": DEFAULT_WINDOW_HEIGHT,
    }

    if not CONFIG_FILE.exists():
        return config

    try:
        loaded_config = json.loads(CONFIG_FILE.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        write_error_log("Gagal membaca file config", exc)
        return config

    if not isinstance(loaded_config, dict):
        return config

    target_path = str(loaded_config.get("target_path") or config["target_path"])
    field_file = str(loaded_config.get("field_file") or config["field_file"])
    if not Path(field_file).exists():
        field_file = str(default_field_file)
    csv_output = str(loaded_config.get("csv_output") or config["csv_output"])
    sheet_mode = str(loaded_config.get("sheet_mode") or "auto").strip().lower()
    if sheet_mode not in {"auto", "custom"}:
        sheet_mode = "auto"

    sheet_name_raw = loaded_config.get("sheet_name")
    sheet_name = str(sheet_name_raw).strip() if sheet_name_raw else None
    if sheet_mode == "auto":
        sheet_name = None

    follow_up_db_path = str(loaded_config.get("follow_up_db_path") or "").strip()
    if not follow_up_db_path:
        follow_up_db_path = str(FOLLOW_UP_DB_FILE)
    else:
        follow_up_db_path = str(FOLLOW_UP_DB_FILE)
    result_view = str(loaded_config.get("result_view") or "check-results").strip()
    if result_view not in {"check-results", "follow-up-db"}:
        result_view = "check-results"

    return {
        "target_path": target_path,
        "field_file": field_file,
        "sheet_mode": sheet_mode,
        "sheet_name": sheet_name,
        "csv_output": csv_output,
        "follow_up_db_path": follow_up_db_path,
        "result_view": result_view,
        "window_width": sanitize_window_size(
            loaded_config.get("window_width"), DEFAULT_WINDOW_WIDTH
        ),
        "window_height": sanitize_window_size(
            loaded_config.get("window_height"), DEFAULT_WINDOW_HEIGHT
        ),
    }


def resolve_initial_directory(path_value: str | None, fallback: Path) -> str:
    if not path_value:
        return str(fallback)

    candidate = Path(path_value)
    if not candidate.is_absolute():
        candidate = (APP_DIR / candidate).resolve()

    if candidate.is_file():
        candidate = candidate.parent

    if candidate.exists():
        return str(candidate)

    if candidate.parent.exists():
        return str(candidate.parent)

    return str(fallback)


def with_hidden_tk_root(action):
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        return action(root)
    finally:
        root.destroy()


FOLLOW_UP_DB_REQUIRED_COLUMNS = (
    "nama_file",
    "countermeasure",
    "responsible",
    "due_date",
    "status",
)
FOLLOW_UP_DB_FIELDNAMES = (*FOLLOW_UP_DB_REQUIRED_COLUMNS, "standard_id")


def normalize_follow_up_db_column_name(value: object) -> str:
    text = str(value or "").strip().casefold().replace("_", " ")
    return " ".join(text.split())


def map_follow_up_db_headers(headers: list[object]) -> dict[str, int]:
    aliases = {
        "nama file": "nama_file",
        "file name": "nama_file",
        "filename": "nama_file",
        "countermeasure": "countermeasure",
        "responsible": "responsible",
        "due date": "due_date",
        "duedate": "due_date",
        "status": "status",
        "standard id": "standard_id",
        "standardid": "standard_id",
    }
    mapped_columns: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized_header = normalize_follow_up_db_column_name(header)
        target_name = aliases.get(normalized_header)
        if target_name and target_name not in mapped_columns:
            mapped_columns[target_name] = index

    missing_columns = [
        column_name
        for column_name in FOLLOW_UP_DB_REQUIRED_COLUMNS
        if column_name not in mapped_columns
    ]
    if missing_columns:
        missing_label = ", ".join(missing_columns)
        raise ValueError(f"Kolom database follow up belum lengkap: {missing_label}")

    return mapped_columns


def read_follow_up_database_rows(file_path: Path) -> list[dict[str, str]]:
    suffix = file_path.suffix.lower()
    rows: list[dict[str, str]] = []

    if suffix == ".csv":
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError("File CSV follow up tidak memiliki header.")
            mapped_columns = map_follow_up_db_headers(list(reader.fieldnames))
            reverse_map = {
                target_name: reader.fieldnames[index]
                for target_name, index in mapped_columns.items()
            }
            optional_standard_id_column = None
            for field_name in reader.fieldnames:
                if normalize_follow_up_db_column_name(field_name) in {
                    "standard id",
                    "standardid",
                }:
                    optional_standard_id_column = field_name
                    break
            for row in reader:
                if not any(str(value or "").strip() for value in row.values()):
                    continue
                rows.append(
                    {
                        "nama_file": str(
                            row.get(reverse_map["nama_file"], "") or ""
                        ).strip(),
                        "countermeasure": str(
                            row.get(reverse_map["countermeasure"], "") or ""
                        ).strip(),
                        "responsible": str(
                            row.get(reverse_map["responsible"], "") or ""
                        ).strip(),
                        "due_date": str(
                            row.get(reverse_map["due_date"], "") or ""
                        ).strip(),
                        "status": str(row.get(reverse_map["status"], "") or "").strip(),
                        "standard_id": str(
                            row.get(optional_standard_id_column, "") or ""
                        ).strip(),
                    }
                )
            return rows

    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            worksheet_rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not worksheet_rows:
            raise ValueError("File Excel follow up kosong.")

        mapped_columns = map_follow_up_db_headers(list(worksheet_rows[0]))
        optional_standard_id_index = None
        for index, header in enumerate(worksheet_rows[0]):
            if normalize_follow_up_db_column_name(header) in {
                "standard id",
                "standardid",
            }:
                optional_standard_id_index = index
                break
        for raw_row in worksheet_rows[1:]:
            if not any(raw_row):
                continue
            rows.append(
                {
                    "nama_file": str(
                        raw_row[mapped_columns["nama_file"]] or ""
                    ).strip(),
                    "countermeasure": str(
                        raw_row[mapped_columns["countermeasure"]] or ""
                    ).strip(),
                    "responsible": str(
                        raw_row[mapped_columns["responsible"]] or ""
                    ).strip(),
                    "due_date": str(raw_row[mapped_columns["due_date"]] or "").strip(),
                    "status": str(raw_row[mapped_columns["status"]] or "").strip(),
                    "standard_id": str(
                        raw_row[optional_standard_id_index] or ""
                    ).strip()
                    if optional_standard_id_index is not None
                    and optional_standard_id_index < len(raw_row)
                    else "",
                }
            )
        return rows

    raise ValueError("Format database follow up harus CSV atau XLSX.")


def normalize_follow_up_database_row(row: dict[str, object]) -> dict[str, str]:
    return {
        "nama_file": str(row.get("nama_file") or "").strip(),
        "countermeasure": str(row.get("countermeasure") or "").strip(),
        "responsible": str(row.get("responsible") or "").strip(),
        "due_date": str(row.get("due_date") or "").strip(),
        "status": str(row.get("status") or "").strip(),
        "standard_id": str(row.get("standard_id") or "").strip(),
    }


def build_follow_up_database_key(row: dict[str, object]) -> tuple[str, str, str, str]:
    normalized_row = normalize_follow_up_database_row(row)
    return (
        normalized_row["nama_file"].casefold(),
        normalized_row["countermeasure"].casefold(),
        normalized_row["responsible"].casefold(),
        normalized_row["due_date"].casefold(),
    )


def ensure_follow_up_database_file(file_path: Path) -> None:
    if file_path.exists():
        return
    file_path.parent.mkdir(parents=True, exist_ok=True)
    with file_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(FOLLOW_UP_DB_FIELDNAMES))
        writer.writeheader()


def write_follow_up_database_rows(
    file_path: Path, rows: list[dict[str, object]]
) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    normalized_rows = [normalize_follow_up_database_row(row) for row in rows]
    with file_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(FOLLOW_UP_DB_FIELDNAMES))
        writer.writeheader()
        writer.writerows(normalized_rows)


def merge_follow_up_database_rows(
    existing_rows: list[dict[str, str]], results: list[dict[str, object]]
) -> list[dict[str, str]]:
    merged_by_key: dict[tuple[str, str, str, str], dict[str, str]] = {}

    for row in existing_rows:
        normalized_row = normalize_follow_up_database_row(row)
        if not any(normalized_row.values()):
            continue
        merged_by_key[build_follow_up_database_key(normalized_row)] = normalized_row

    for result in results:
        file_name = str(result.get("file_name") or "").strip()
        follow_up_rows = result.get("follow_up_rows") or []
        for follow_up_row in follow_up_rows:
            candidate_row = normalize_follow_up_database_row(
                {
                    "nama_file": file_name,
                    "countermeasure": follow_up_row.get("countermeasure"),
                    "responsible": follow_up_row.get("responsible"),
                    "due_date": follow_up_row.get("due_date"),
                    "status": "Open",
                    "standard_id": "",
                }
            )
            if not any(
                (
                    candidate_row["countermeasure"],
                    candidate_row["responsible"],
                    candidate_row["due_date"],
                )
            ):
                continue

            row_key = build_follow_up_database_key(candidate_row)
            existing_row = merged_by_key.get(row_key)
            if existing_row and existing_row.get("status"):
                candidate_row["status"] = existing_row["status"]
                candidate_row["standard_id"] = existing_row.get("standard_id", "")
            merged_by_key[row_key] = candidate_row

    return sorted(
        merged_by_key.values(),
        key=lambda row: (
            row["nama_file"].casefold(),
            row["due_date"].casefold(),
            row["countermeasure"].casefold(),
            row["responsible"].casefold(),
        ),
    )


def main(page: ft.Page) -> None:
    try:
        initialize_app_storage()
    except Exception as exc:
        write_error_log("Gagal menyiapkan folder data aplikasi", exc)

    section_labels = {
        "1.1": "1.1. Initiate IPS",
        "1.2": "1.2. Refocus the Problem",
        "1.3": "1.3. Verify base condition",
        "1.4": "1.4. Restore Base Condition",
        "1.5": "1.5. Ask WHY WHY",
    }
    side_button_width = 190

    default_field_file = get_default_field_file_path()

    default_target_path = APP_DIR / "new create ips"
    if not default_target_path.exists():
        default_target_path = APP_DIR

    persisted_config = load_app_config(default_target_path, default_field_file)

    page.title = "IPS Completeness Checker"
    page.window_width = int(persisted_config["window_width"])
    page.window_height = int(persisted_config["window_height"])
    page.padding = 24
    page.theme_mode = ft.ThemeMode.LIGHT
    page.bgcolor = "#f4f1e8"
    page.scroll = ft.ScrollMode.HIDDEN
    if ICON_FILE.exists():
        page.window.icon = str(ICON_FILE)
    page.theme = ft.Theme(
        color_scheme=ft.ColorScheme(
            primary="#0f5c4d",
            secondary="#c96f3b",
            surface="#fffaf0",
        ),
        visual_density=ft.VisualDensity.COMFORTABLE,
    )

    summary_text = ft.Text(size=16, weight=ft.FontWeight.W_600, color="#16302b")
    loading_snack_bar: ft.SnackBar | None = None
    app_state = {
        "target_path": str(persisted_config["target_path"]),
        "field_file": str(persisted_config["field_file"]),
        "sheet_mode": str(persisted_config["sheet_mode"]),
        "sheet_name": persisted_config["sheet_name"],
        "csv_output": str(persisted_config["csv_output"]),
        "follow_up_db_path": str(persisted_config.get("follow_up_db_path") or ""),
        "result_view": str(persisted_config.get("result_view") or "check-results"),
        "window_width": int(persisted_config["window_width"]),
        "window_height": int(persisted_config["window_height"]),
        "follow_up_db_rows": [],
        "follow_up_db_error": "",
        "is_loading": False,
        "is_picker_open": False,
    }
    table_state = {
        "results": [],
        "sort_column_index": 1,
        "sort_ascending": True,
        "follow_up_sort_column_index": 0,
        "follow_up_sort_ascending": True,
    }
    results_table = ft.Column(
        spacing=0,
        expand=True,
        scroll=ft.ScrollMode.AUTO,
    )
    table_container = ft.Container(
        expand=True,
        padding=0,
        margin=0,
        bgcolor="#fffaf0",
        border=ft.Border.all(1, "#d9c9ab"),
        border_radius=18,
        clip_behavior=ft.ClipBehavior.HARD_EDGE,
        content=results_table,
    )
    follow_up_db_table = ft.Column(
        spacing=0,
        expand=True,
        scroll=ft.ScrollMode.AUTO,
    )
    follow_up_db_table_container = ft.Container(
        expand=True,
        padding=0,
        margin=0,
        bgcolor="#fffaf0",
        border=ft.Border.all(1, "#d9c9ab"),
        border_radius=18,
        clip_behavior=ft.ClipBehavior.HARD_EDGE,
        content=follow_up_db_table,
    )
    result_panel_title = ft.Text(
        "Hasil Pemeriksaan",
        size=24,
        weight=ft.FontWeight.W_700,
        color="#16302b",
    )
    result_view_menu_button = ft.PopupMenuButton(
        icon=ft.Icons.MENU,
        tooltip="Pilih tampilan panel kanan",
        style=ft.ButtonStyle(
            bgcolor="#e9dfc9",
            color="#16302b",
            padding=ft.Padding.symmetric(horizontal=8, vertical=8),
            shape=ft.RoundedRectangleBorder(radius=12),
            side=ft.BorderSide(1, "#d9c9ab"),
        ),
        items=[],
    )
    export_action_button = ft.Button(
        content=ft.Text("Export", weight=ft.FontWeight.W_600),
        icon=ft.Icons.DOWNLOAD_ROUNDED,
        width=156,
        style=ft.ButtonStyle(
            bgcolor="#0f5c4d",
            color="#fffaf0",
            padding=ft.Padding.symmetric(horizontal=16, vertical=12),
            shape=ft.RoundedRectangleBorder(radius=14),
        ),
    )
    result_view_menu = ft.AppBar(
        bgcolor="#fffaf0",
        toolbar_height=64,
        leading=ft.Container(
            content=result_view_menu_button,
            padding=ft.Padding.all(12),
        ),
        leading_width=64,
        title=ft.Column(
            spacing=2,
            tight=True,
            controls=[
                result_panel_title,
            ],
        ),
        center_title=False,
        elevation=0,
        actions=[
            ft.Container(
                padding=ft.Padding.only(right=12),
                content=export_action_button,
            )
        ],
    )
    result_view_content = ft.Container(expand=True)
    detail_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text(
            "Detail Completeness IPS", weight=ft.FontWeight.W_700, color="#16302b"
        ),
        actions=[ft.TextButton("Tutup")],
        actions_alignment=ft.MainAxisAlignment.END,
        scrollable=True,
    )
    follow_up_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text(
            "Follow Up Section 1.5", weight=ft.FontWeight.W_700, color="#16302b"
        ),
        actions=[ft.TextButton("Tutup")],
        actions_alignment=ft.MainAxisAlignment.END,
        scrollable=True,
    )
    share_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text(
            "Share Detail Completeness", weight=ft.FontWeight.W_700, color="#16302b"
        ),
        actions=[ft.TextButton("Tutup")],
        actions_alignment=ft.MainAxisAlignment.END,
        scrollable=True,
    )
    edit_follow_up_status_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text(
            "Edit Status Follow Up", weight=ft.FontWeight.W_700, color="#16302b"
        ),
        actions_alignment=ft.MainAxisAlignment.END,
    )
    delete_follow_up_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text(
            "Hapus Data Follow Up", weight=ft.FontWeight.W_700, color="#16302b"
        ),
        actions_alignment=ft.MainAxisAlignment.END,
    )
    edit_follow_up_status_options = ["Open", "Close"]
    edit_follow_up_status_field = ft.Dropdown(
        label="Status",
        border_color="#c6b89b",
        focused_border_color="#0f5c4d",
        text_style=ft.TextStyle(color="#16302b"),
        label_style=ft.TextStyle(color="#6b4f3a"),
        options=[
            ft.dropdown.Option(option) for option in edit_follow_up_status_options
        ],
        value="Open",
        filled=True,
        bgcolor="#fffaf0",
    )
    edit_follow_up_standard_id_field = ft.TextField(
        label="Standard ID",
        border_color="#c6b89b",
        focused_border_color="#0f5c4d",
        text_style=ft.TextStyle(color="#16302b"),
        label_style=ft.TextStyle(color="#6b4f3a"),
        expand=True,
    )
    edit_follow_up_status_hint = ft.Text(color="#6b4f3a", size=12)
    edit_follow_up_standard_id_hint = ft.Text(color="#6b4f3a", size=12)
    edit_follow_up_status_target: dict[str, str] = {}
    delete_follow_up_target: dict[str, str] = {}
    info_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text("Info GUI", weight=ft.FontWeight.W_700, color="#16302b"),
        actions=[ft.TextButton("Tutup")],
        actions_alignment=ft.MainAxisAlignment.END,
        scrollable=True,
    )
    export_dialog = ft.AlertDialog(
        modal=True,
        bgcolor="#fffaf0",
        shape=ft.RoundedRectangleBorder(radius=18),
        title=ft.Text("Pilih Tipe Export", weight=ft.FontWeight.W_700, color="#16302b"),
        actions=[ft.TextButton("Tutup")],
        actions_alignment=ft.MainAxisAlignment.END,
        scrollable=True,
    )
    clipboard_service = ft.Clipboard()
    folder_button = ft.Button(
        content=ft.Text("Pilih Folder", weight=ft.FontWeight.W_600),
        icon=ft.Icons.FOLDER_OPEN_ROUNDED,
        width=side_button_width,
        style=ft.ButtonStyle(
            bgcolor="#e9dfc9",
            color="#16302b",
            padding=ft.Padding.symmetric(horizontal=18, vertical=16),
            shape=ft.RoundedRectangleBorder(radius=14),
        ),
    )
    file_button = ft.Button(
        content=ft.Text("Pilih File", weight=ft.FontWeight.W_600),
        icon=ft.Icons.UPLOAD_FILE_ROUNDED,
        width=side_button_width,
        style=ft.ButtonStyle(
            bgcolor="#e9dfc9",
            color="#16302b",
            padding=ft.Padding.symmetric(horizontal=18, vertical=16),
            shape=ft.RoundedRectangleBorder(radius=14),
        ),
    )
    create_ips_button = ft.Button(
        content=ft.Text("Buat IPS", weight=ft.FontWeight.W_600),
        icon=ft.Icons.DESCRIPTION_ROUNDED,
        width=side_button_width,
        style=ft.ButtonStyle(
            bgcolor="#e9dfc9",
            color="#16302b",
            padding=ft.Padding.symmetric(horizontal=18, vertical=16),
            shape=ft.RoundedRectangleBorder(radius=14),
        ),
    )
    info_button = ft.Button(
        content=ft.Text("Info", weight=ft.FontWeight.W_600),
        icon=ft.Icons.INFO_OUTLINE_ROUNDED,
        width=side_button_width,
        style=ft.ButtonStyle(
            bgcolor="#e9dfc9",
            color="#16302b",
            padding=ft.Padding.symmetric(horizontal=18, vertical=16),
            shape=ft.RoundedRectangleBorder(radius=14),
        ),
    )
    app_logo = ft.Container(
        width=side_button_width,
        padding=ft.Padding.symmetric(horizontal=14, vertical=14),
        bgcolor="#efe4ce",
        border=ft.Border.all(1, "#d9c9ab"),
        border_radius=18,
        content=ft.Column(
            spacing=10,
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            controls=[
                ft.Container(
                    width=54,
                    height=54,
                    bgcolor="#0f5c4d",
                    border_radius=16,
                    content=ft.Stack(
                        controls=[
                            ft.Container(
                                left=0,
                                right=0,
                                top=0,
                                bottom=0,
                                content=ft.Text(
                                    "IPS",
                                    color="#fffaf0",
                                    weight=ft.FontWeight.W_700,
                                    text_align=ft.TextAlign.CENTER,
                                    size=15,
                                ),
                            ),
                            ft.Container(
                                right=4,
                                bottom=4,
                                width=14,
                                height=14,
                                bgcolor="#c96f3b",
                                border_radius=999,
                            ),
                        ]
                    ),
                ),
                ft.Column(
                    spacing=2,
                    tight=True,
                    horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                    controls=[
                        ft.Text(
                            "IPS Checker",
                            weight=ft.FontWeight.W_700,
                            color="#16302b",
                            size=17,
                            text_align=ft.TextAlign.CENTER,
                        ),
                        ft.Text(
                            "Monitor completeness form IPS",
                            color="#6b4f3a",
                            size=11,
                            text_align=ft.TextAlign.CENTER,
                        ),
                    ],
                ),
            ],
        ),
    )

    def refresh_busy_state() -> None:
        is_loading = bool(app_state["is_loading"])
        is_picker_open = bool(app_state["is_picker_open"])
        is_busy = is_loading or is_picker_open
        folder_button.disabled = is_busy
        file_button.disabled = is_busy
        create_ips_button.disabled = is_busy
        export_action_button.disabled = is_busy
        info_button.disabled = is_busy

    def show_loading_snack_bar() -> None:
        nonlocal loading_snack_bar
        if loading_snack_bar and getattr(loading_snack_bar, "open", False):
            return
        loading_snack_bar = ft.SnackBar(
            bgcolor="#245a4a",
            duration=600000,
            open=True,
            content=ft.Row(
                tight=True,
                spacing=12,
                controls=[
                    ft.ProgressRing(
                        width=18,
                        height=18,
                        stroke_width=2.5,
                        color="#fffaf0",
                    ),
                    ft.Text(
                        "Sedang memeriksa file Excel...",
                        color="#fffaf0",
                        weight=ft.FontWeight.W_600,
                    ),
                ],
            ),
        )
        page.show_dialog(loading_snack_bar)

    def hide_loading_snack_bar() -> None:
        nonlocal loading_snack_bar
        if loading_snack_bar and getattr(loading_snack_bar, "open", False):
            loading_snack_bar.open = False
            loading_snack_bar.update()
        loading_snack_bar = None

    def persist_config() -> None:
        config_payload = {
            "target_path": app_state["target_path"],
            "field_file": app_state["field_file"],
            "sheet_mode": app_state["sheet_mode"],
            "sheet_name": app_state["sheet_name"],
            "csv_output": app_state["csv_output"],
            "follow_up_db_path": app_state["follow_up_db_path"],
            "result_view": app_state["result_view"],
            "window_width": app_state["window_width"],
            "window_height": app_state["window_height"],
        }
        try:
            CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
            CONFIG_FILE.write_text(
                json.dumps(config_payload, indent=2), encoding="utf-8"
            )
        except Exception as exc:
            write_error_log("Gagal menyimpan file config", exc)

    def sync_window_size_from_page() -> None:
        width = page.window_width or page.width or app_state["window_width"]
        height = page.window_height or page.height or app_state["window_height"]
        app_state["window_width"] = sanitize_window_size(
            width, int(app_state["window_width"])
        )
        app_state["window_height"] = sanitize_window_size(
            height, int(app_state["window_height"])
        )

    def set_loading(is_loading: bool) -> None:
        app_state["is_loading"] = is_loading
        refresh_busy_state()
        if is_loading:
            show_loading_snack_bar()
            return
        hide_loading_snack_bar()

    def set_picker_open(is_picker_open: bool) -> None:
        app_state["is_picker_open"] = is_picker_open
        refresh_busy_state()

    def notify(message: str, error: bool = False) -> None:
        page.show_dialog(
            ft.SnackBar(
                content=ft.Text(message),
                bgcolor="#a63d40" if error else "#245a4a",
                open=True,
            )
        )

    def build_result_view_menu_item(label: str, view_name: str) -> ft.PopupMenuItem:
        is_active = str(app_state["result_view"]) == view_name
        return ft.PopupMenuItem(
            content=ft.Row(
                spacing=10,
                controls=[
                    ft.Icon(
                        ft.Icons.CHECK_ROUNDED
                        if is_active
                        else ft.Icons.CIRCLE_OUTLINED,
                        color="#0f5c4d" if is_active else "#9a8874",
                        size=16,
                    ),
                    ft.Text(
                        label,
                        color="#16302b",
                        weight=ft.FontWeight.W_700
                        if is_active
                        else ft.FontWeight.W_500,
                    ),
                ],
            ),
            on_click=lambda _, target=view_name: set_result_view(target),
        )

    def refresh_result_view_menu() -> None:
        current_view = str(app_state["result_view"])
        result_view_menu_button.items = [
            build_result_view_menu_item("Tabel Hasil Pemeriksaan", "check-results"),
            build_result_view_menu_item(
                "Tabel Follow Up Countermeasure", "follow-up-db"
            ),
        ]
        if current_view == "follow-up-db":
            result_view_menu_button.icon_color = "#c96f3b"
            export_action_button.content = ft.Text("Export", weight=ft.FontWeight.W_600)
            export_action_button.style = ft.ButtonStyle(
                bgcolor="#c96f3b",
                color="#fffaf0",
                padding=ft.Padding.symmetric(horizontal=16, vertical=12),
                shape=ft.RoundedRectangleBorder(radius=14),
            )
            return

        result_view_menu_button.icon_color = "#0f5c4d"
        export_action_button.content = ft.Text("Export", weight=ft.FontWeight.W_600)
        export_action_button.style = ft.ButtonStyle(
            bgcolor="#0f5c4d",
            color="#fffaf0",
            padding=ft.Padding.symmetric(horizontal=16, vertical=12),
            shape=ft.RoundedRectangleBorder(radius=14),
        )

    def refresh_follow_up_db_table() -> None:
        header_row = ft.Container(
            bgcolor="#efe4ce",
            padding=ft.Padding.symmetric(horizontal=12, vertical=14),
            content=ft.Row(
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                controls=[
                    ft.Container(
                        width=220,
                        content=build_sortable_header_cell(
                            "Nama File",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            0,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        expand=2,
                        content=build_sortable_header_cell(
                            "Countermeasure",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            1,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        width=115,
                        content=build_sortable_header_cell(
                            "Responsible",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            2,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        width=95,
                        content=build_sortable_header_cell(
                            "Due Date",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            3,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        width=78,
                        content=build_sortable_header_cell(
                            "Status",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            4,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        width=96,
                        content=build_sortable_header_cell(
                            "Standard ID",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            5,
                            toggle_follow_up_sort,
                        ),
                    ),
                    ft.Container(
                        width=114,
                        content=build_sortable_header_cell(
                            "Action",
                            table_state["follow_up_sort_column_index"],
                            table_state["follow_up_sort_ascending"],
                            None,
                            toggle_follow_up_sort,
                        ),
                    ),
                ],
            ),
        )

        rows = list(app_state["follow_up_db_rows"])
        result_rows: list[ft.Control] = []
        for index, row in enumerate(rows):
            status_value = str(row.get("status") or "-")
            standard_id_value = str(row.get("standard_id") or "-")
            row_bg = "#fffaf0" if index % 2 == 0 else "#fcf5e8"
            if status_value.casefold() in {"done", "complete", "close", "closed"}:
                status_color = "#245a4a"
            elif status_value.casefold() in {
                "open",
                "ongoing",
                "progress",
                "in progress",
            }:
                status_color = "#c96f3b"
            else:
                status_color = "#6b4f3a"
            result_rows.append(
                ft.Container(
                    bgcolor=row_bg,
                    border=ft.Border(bottom=ft.BorderSide(1, "#eadfcb")),
                    padding=ft.Padding.symmetric(horizontal=12, vertical=14),
                    content=ft.Row(
                        vertical_alignment=ft.CrossAxisAlignment.START,
                        controls=[
                            ft.Container(
                                width=220,
                                content=ft.Text(
                                    str(row.get("nama_file") or "-"),
                                    color="#16302b",
                                    max_lines=2,
                                    overflow=ft.TextOverflow.ELLIPSIS,
                                ),
                            ),
                            ft.Container(
                                expand=2,
                                content=ft.Text(
                                    str(row.get("countermeasure") or "-"),
                                    color="#16302b",
                                ),
                            ),
                            ft.Container(
                                width=115,
                                content=ft.Text(
                                    str(row.get("responsible") or "-"),
                                    color="#4b5d58",
                                    max_lines=2,
                                    overflow=ft.TextOverflow.ELLIPSIS,
                                ),
                            ),
                            ft.Container(
                                width=95,
                                content=ft.Text(
                                    str(row.get("due_date") or "-"),
                                    color="#4b5d58",
                                    max_lines=2,
                                    overflow=ft.TextOverflow.ELLIPSIS,
                                ),
                            ),
                            ft.Container(
                                width=78,
                                content=ft.Text(
                                    status_value,
                                    color=status_color,
                                    weight=ft.FontWeight.W_600,
                                    max_lines=2,
                                    overflow=ft.TextOverflow.ELLIPSIS,
                                ),
                            ),
                            ft.Container(
                                width=96,
                                content=ft.Text(
                                    standard_id_value,
                                    color="#4b5d58",
                                    max_lines=2,
                                    overflow=ft.TextOverflow.ELLIPSIS,
                                ),
                            ),
                            ft.Container(
                                width=114,
                                content=ft.Row(
                                    spacing=2,
                                    alignment=ft.MainAxisAlignment.CENTER,
                                    controls=[
                                        ft.IconButton(
                                            icon=ft.Icons.ASSIGNMENT_TURNED_IN_ROUNDED,
                                            icon_color="#0f5c4d",
                                            icon_size=16,
                                            style=ft.ButtonStyle(padding=0),
                                            tooltip="Lihat follow up",
                                            on_click=lambda _, item=row: (
                                                show_follow_up_details_from_database_row(
                                                    item
                                                )
                                            ),
                                        ),
                                        ft.IconButton(
                                            icon=ft.Icons.EDIT_NOTE_ROUNDED,
                                            icon_color="#0f5c4d",
                                            icon_size=16,
                                            style=ft.ButtonStyle(padding=0),
                                            tooltip="Edit status",
                                            on_click=lambda _, item=row: (
                                                show_edit_follow_up_status_dialog(item)
                                            ),
                                        ),
                                        ft.IconButton(
                                            icon=ft.Icons.DELETE_OUTLINE_ROUNDED,
                                            icon_color="#a63d40",
                                            icon_size=16,
                                            style=ft.ButtonStyle(padding=0),
                                            tooltip="Hapus data",
                                            on_click=lambda _, item=row: (
                                                show_delete_follow_up_dialog(item)
                                            ),
                                        ),
                                    ],
                                ),
                            ),
                        ],
                    ),
                )
            )

        if not result_rows:
            empty_message = (
                app_state["follow_up_db_error"]
                or "Belum ada data follow up pada database lokal otomatis."
            )
            result_rows = [
                ft.Container(
                    padding=24,
                    content=ft.Text(str(empty_message), color="#6b4f3a"),
                )
            ]

        follow_up_db_table.controls = [header_row, *result_rows]

    def refresh_result_panel_view() -> None:
        refresh_result_view_menu()
        current_view = str(app_state["result_view"])
        if current_view == "follow-up-db":
            result_panel_title.value = "Follow Up Countermeasure"
            result_view_content.content = ft.Column(
                expand=True,
                spacing=16,
                controls=[
                    follow_up_db_table_container,
                ],
            )
            return

        result_panel_title.value = "Hasil Pemeriksaan"
        result_view_content.content = ft.Column(
            expand=True,
            spacing=16,
            controls=[
                table_container,
            ],
        )

    def set_result_view(view_name: str) -> None:
        if view_name not in {"check-results", "follow-up-db"}:
            return
        app_state["result_view"] = view_name
        persist_config()
        refresh_result_panel_view()
        page.update()

    def load_follow_up_database(update_page: bool = True) -> None:
        database_file = FOLLOW_UP_DB_FILE
        app_state["follow_up_db_path"] = str(database_file)
        try:
            ensure_follow_up_database_file(database_file)
            rows = read_follow_up_database_rows(database_file)
        except Exception as exc:
            app_state["follow_up_db_rows"] = []
            app_state["follow_up_db_error"] = str(exc)
            persist_config()
            refresh_follow_up_db_table()
            if update_page:
                notify(str(exc), error=True)
                page.update()
            return

        app_state["follow_up_db_rows"] = rows
        app_state["follow_up_db_error"] = ""
        sort_follow_up_rows(
            table_state["follow_up_sort_column_index"],
            table_state["follow_up_sort_ascending"],
        )
        persist_config()
        if update_page:
            notify("Database follow up otomatis berhasil dimuat.")
            page.update()

    def sync_follow_up_database_from_results(
        results: list[dict[str, object]], update_page: bool = False
    ) -> None:
        database_file = FOLLOW_UP_DB_FILE
        app_state["follow_up_db_path"] = str(database_file)
        try:
            ensure_follow_up_database_file(database_file)
            existing_rows = read_follow_up_database_rows(database_file)
            merged_rows = merge_follow_up_database_rows(existing_rows, results)
            write_follow_up_database_rows(database_file, merged_rows)
        except Exception as exc:
            app_state["follow_up_db_error"] = str(exc)
            refresh_follow_up_db_table()
            if update_page:
                log_path = write_error_log("Sinkron database follow up gagal", exc)
                notify(str(exc), error=True)
                notify(f"Log error tersimpan di: {log_path.name}", error=True)
                page.update()
            return

        app_state["follow_up_db_rows"] = merged_rows
        app_state["follow_up_db_error"] = ""
        sort_follow_up_rows(
            table_state["follow_up_sort_column_index"],
            table_state["follow_up_sort_ascending"],
        )
        persist_config()
        if update_page:
            page.update()

    def pick_target_folder(_: ft.ControlEvent) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return
        set_picker_open(True)
        page.update()
        selected_path = with_hidden_tk_root(
            lambda _: filedialog.askdirectory(
                title="Pilih folder Excel",
                initialdir=resolve_initial_directory(
                    str(app_state["target_path"]), APP_DIR
                ),
            )
        )
        if selected_path:
            app_state["target_path"] = selected_path
            persist_config()
            set_picker_open(False)
            start_run_check(selected_path)
            return
        set_picker_open(False)
        page.update()

    def pick_target_file(_: ft.ControlEvent) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return
        set_picker_open(True)
        page.update()
        selected_path = with_hidden_tk_root(
            lambda _: filedialog.askopenfilename(
                title="Pilih file Excel",
                initialdir=resolve_initial_directory(
                    str(app_state["target_path"]), APP_DIR
                ),
                filetypes=[("Excel files", "*.xlsx")],
            )
        )
        if selected_path:
            app_state["target_path"] = selected_path
            persist_config()
            set_picker_open(False)
            start_run_check(selected_path)
            return
        set_picker_open(False)
        page.update()

    def handle_page_resize(_: ft.ControlEvent) -> None:
        sync_window_size_from_page()
        refresh_layout_height()
        persist_config()
        page.update()

    def handle_page_close(_: ft.ControlEvent) -> None:
        sync_window_size_from_page()
        refresh_layout_height()
        persist_config()

    def build_completeness_cell(result: dict[str, object]) -> ft.Control:
        completeness = float(result["completeness_percentage"])
        if completeness >= 80:
            bar_color = "#245a4a"
        elif completeness >= 50:
            bar_color = "#c96f3b"
        else:
            bar_color = "#a63d40"

        return ft.Row(
            spacing=10,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
            controls=[
                ft.ProgressBar(
                    value=completeness / 100,
                    width=120,
                    bar_height=8,
                    color=bar_color,
                    bgcolor="#e7dbc5",
                    border_radius=999,
                ),
                ft.Text(
                    f"{completeness:.2f}%", weight=ft.FontWeight.W_600, color="#16302b"
                ),
            ],
        )

    def close_detail_dialog(_: ft.ControlEvent) -> None:
        detail_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_follow_up_dialog(_: ft.ControlEvent) -> None:
        follow_up_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_share_dialog(_: ft.ControlEvent) -> None:
        share_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_info_dialog(_: ft.ControlEvent) -> None:
        info_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_export_dialog(_: ft.ControlEvent) -> None:
        export_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_edit_follow_up_status_dialog(_: ft.ControlEvent | None = None) -> None:
        edit_follow_up_status_dialog.open = False
        page.pop_dialog()
        page.update()

    def close_delete_follow_up_dialog(_: ft.ControlEvent | None = None) -> None:
        delete_follow_up_dialog.open = False
        page.pop_dialog()
        page.update()

    async def copy_summary_to_clipboard(summary_text_value: str) -> None:
        try:
            await clipboard_service.set(summary_text_value)
            notify("Summary copied to clipboard.")
        except Exception as exc:
            notify(f"Gagal menyalin summary: {exc}", error=True)

    def save_follow_up_status_change(_: ft.ControlEvent) -> None:
        updated_status = str(edit_follow_up_status_field.value or "").strip()
        updated_standard_id = str(edit_follow_up_standard_id_field.value or "").strip()
        if not updated_status:
            notify("Status follow up tidak boleh kosong.", error=True)
            return
        if updated_status == "Close" and not updated_standard_id:
            notify(
                "Standard ID wajib diisi jika ingin mengubah status menjadi Close.",
                error=True,
            )
            return

        target_key = build_follow_up_database_key(edit_follow_up_status_target)
        updated_rows: list[dict[str, str]] = []
        row_found = False
        for row in app_state["follow_up_db_rows"]:
            normalized_row = normalize_follow_up_database_row(row)
            if build_follow_up_database_key(normalized_row) == target_key:
                normalized_row["status"] = updated_status
                normalized_row["standard_id"] = (
                    updated_standard_id if updated_status == "Close" else ""
                )
                row_found = True
            updated_rows.append(normalized_row)

        if not row_found:
            notify("Data follow up yang dipilih tidak ditemukan.", error=True)
            return

        try:
            write_follow_up_database_rows(FOLLOW_UP_DB_FILE, updated_rows)
        except Exception as exc:
            log_path = write_error_log("Simpan status follow up gagal", exc)
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        app_state["follow_up_db_rows"] = updated_rows
        app_state["follow_up_db_error"] = ""
        sort_follow_up_rows(
            table_state["follow_up_sort_column_index"],
            table_state["follow_up_sort_ascending"],
        )
        close_edit_follow_up_status_dialog(None)
        notify("Status follow up berhasil diperbarui.")

    def delete_follow_up_row(_: ft.ControlEvent) -> None:
        target_key = build_follow_up_database_key(delete_follow_up_target)
        updated_rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
            if build_follow_up_database_key(normalize_follow_up_database_row(row))
            != target_key
        ]

        if len(updated_rows) == len(app_state["follow_up_db_rows"]):
            notify("Data follow up yang dipilih tidak ditemukan.", error=True)
            return

        try:
            write_follow_up_database_rows(FOLLOW_UP_DB_FILE, updated_rows)
        except Exception as exc:
            log_path = write_error_log("Hapus data follow up gagal", exc)
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        app_state["follow_up_db_rows"] = updated_rows
        app_state["follow_up_db_error"] = ""
        sort_follow_up_rows(
            table_state["follow_up_sort_column_index"],
            table_state["follow_up_sort_ascending"],
        )
        close_delete_follow_up_dialog(None)
        notify("Data follow up berhasil dihapus.")

    def show_edit_follow_up_status_dialog(row: dict[str, object]) -> None:
        normalized_row = normalize_follow_up_database_row(row)
        edit_follow_up_status_target.clear()
        edit_follow_up_status_target.update(normalized_row)
        current_status = normalized_row.get("status") or "Open"
        if current_status.casefold() == "closed":
            current_status = "Close"
        elif current_status != "Close":
            current_status = "Open"
        edit_follow_up_status_field.options = [
            ft.dropdown.Option(option) for option in edit_follow_up_status_options
        ]
        edit_follow_up_status_field.value = current_status
        edit_follow_up_standard_id_field.disabled = False
        edit_follow_up_standard_id_field.read_only = False
        edit_follow_up_standard_id_field.value = normalized_row.get("standard_id") or ""
        edit_follow_up_status_hint.value = (
            f"{normalized_row.get('nama_file') or '-'} | "
            f"{normalized_row.get('countermeasure') or '-'}"
        )
        edit_follow_up_standard_id_hint.value = (
            "Isi Standard ID jika status dipilih Close."
        )
        edit_follow_up_status_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                spacing=12,
                tight=True,
                controls=[
                    edit_follow_up_status_hint,
                    ft.Row(
                        spacing=12,
                        controls=[
                            ft.Container(
                                width=180, content=edit_follow_up_status_field
                            ),
                            edit_follow_up_standard_id_field,
                        ],
                    ),
                    edit_follow_up_standard_id_hint,
                ],
            ),
        )
        edit_follow_up_status_dialog.actions = [
            ft.TextButton("Simpan", on_click=save_follow_up_status_change),
            ft.TextButton("Tutup", on_click=close_edit_follow_up_status_dialog),
        ]
        if getattr(edit_follow_up_status_dialog, "open", False):
            page.update()
            return
        edit_follow_up_status_dialog.open = True
        page.show_dialog(edit_follow_up_status_dialog)

    def show_delete_follow_up_dialog(row: dict[str, object]) -> None:
        normalized_row = normalize_follow_up_database_row(row)
        delete_follow_up_target.clear()
        delete_follow_up_target.update(normalized_row)
        delete_follow_up_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        "Data follow up ini akan dihapus dari database lokal otomatis.",
                        color="#4b5d58",
                    ),
                    ft.Text(
                        f"{normalized_row.get('nama_file') or '-'} | {normalized_row.get('countermeasure') or '-'}",
                        color="#6b4f3a",
                        size=12,
                    ),
                ],
            ),
        )
        delete_follow_up_dialog.actions = [
            ft.TextButton("Hapus", on_click=delete_follow_up_row),
            ft.TextButton("Tutup", on_click=close_delete_follow_up_dialog),
        ]
        if getattr(delete_follow_up_dialog, "open", False):
            page.update()
            return
        delete_follow_up_dialog.open = True
        page.show_dialog(delete_follow_up_dialog)

    def build_detail_summary_text(result: dict[str, object]) -> str:
        status_label = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
        section_completeness = result.get("section_completeness", {})
        summary_lines = [
            f"*File: {result['file_name']}*",
            f"> *Participant: {result.get('participant') or '-'}*",
            f"> *Status: {status_label}*",
            f"> *Completeness: {float(result['completeness_percentage']):.2f}%*",
            "> *Per Section:*",
        ]
        for section in sorted(section_completeness):
            percentage = float(section_completeness.get(section, 0.0))
            summary_lines.append(
                f"- {section_labels.get(str(section), str(section))}: {percentage:.1f}%"
            )
        return "\n".join(summary_lines)

    def build_follow_up_summary_text(result: dict[str, object]) -> str:
        follow_up_rows = result.get("follow_up_rows", [])
        summary_lines = [
            f"*File: {result['file_name']}*",
            f"> *Participant: {result.get('participant') or '-'}*",
            "> *Follow Up Section 1.5 - Ask WHY WHY*",
        ]

        if not follow_up_rows:
            summary_lines.append("Belum ada data follow up.")
            return "\n".join(summary_lines)

        for index, row in enumerate(follow_up_rows, start=1):
            summary_lines.extend(
                [
                    f"{index}. Countermeasure: {row.get('countermeasure') or '-'}",
                    f"   Responsible: {row.get('responsible') or '-'}",
                    f"   Due Date: {row.get('due_date') or '-'}",
                ]
            )
        return "\n".join(summary_lines)

    def build_results_summary_text(incomplete_only: bool) -> str:
        if incomplete_only:
            results = [
                result
                for result in table_state["results"]
                if not result.get("is_complete")
            ]
        else:
            results = list(table_state["results"])
        if not results:
            if incomplete_only:
                return "Belum ada file incomplete untuk dibagikan."
            return "Belum ada hasil pemeriksaan untuk dibagikan."

        title = (
            "*Daftar File INCOMPLETE IPS*" if incomplete_only else "*Daftar Data IPS*"
        )
        summary_lines = [
            title,
            f"_{len(results)} File_",
            "",
        ]
        for _index, result in enumerate(results, start=1):
            status_label = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
            completeness = float(result["completeness_percentage"])
            participant = result.get("participant") or "-"
            summary_lines.extend(
                [
                    f"> *{result['file_name']}*",
                    f"- Participant: {participant}",
                    f"- Status: {status_label} ({completeness:.2f}%)\n",
                ]
            )
        return "\n".join(summary_lines)

    def build_follow_up_db_summary_text() -> str:
        rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
        ]
        if not rows:
            return "Belum ada data follow up countermeasure untuk dibagikan."

        close_statuses = {"close", "closed", "complete", "done"}
        close_count = sum(
            1
            for row in rows
            if str(row.get("status") or "").strip().casefold() in close_statuses
        )
        open_count = len(rows) - close_count

        summary_lines = [
            "*Daftar Follow Up Countermeasure*",
            f"_{len(rows)} Data | {open_count} Open | {close_count} Close_",
            "",
        ]
        for row in rows:
            summary_lines.extend(
                [
                    f"> *{row.get('nama_file') or '-'}*",
                    f"- Countermeasure: {row.get('countermeasure') or '-'}",
                    f"- Responsible: {row.get('responsible') or '-'}",
                    f"- Due Date: {row.get('due_date') or '-'}",
                    f"- Status: {row.get('status') or '-'} | Standard ID: {row.get('standard_id') or '-'}\n",
                ]
            )
        return "\n".join(summary_lines)

    def export_follow_up_db_summary_pdf(output_path: Path) -> None:
        rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
        ]
        if not rows:
            raise ValueError(
                "Belum ada data follow up countermeasure untuk dibuatkan PDF."
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=1.2 * cm,
            rightMargin=1.2 * cm,
            topMargin=1.4 * cm,
            bottomMargin=1.2 * cm,
        )
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor("#16302b")
        styles["Normal"].textColor = colors.HexColor("#4b5d58")
        styles["BodyText"].fontSize = 8
        styles["BodyText"].leading = 10

        close_statuses = {"close", "closed", "complete", "done"}
        close_count = sum(
            1
            for row in rows
            if str(row.get("status") or "").strip().casefold() in close_statuses
        )
        open_count = len(rows) - close_count

        story = [
            Paragraph("Follow Up Countermeasure Summary", styles["Title"]),
            Spacer(1, 0.2 * cm),
            Paragraph(
                (
                    f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
                    f"<br/>{len(rows)} data follow up | {open_count} open | {close_count} close"
                ),
                styles["Normal"],
            ),
            Spacer(1, 0.35 * cm),
        ]

        table_rows = [
            [
                "Nama File",
                "Countermeasure",
                "Responsible",
                "Due Date",
                "Status",
                "Standard ID",
            ]
        ]
        for row in rows:
            table_rows.append(
                [
                    Paragraph(str(row.get("nama_file") or "-"), styles["BodyText"]),
                    Paragraph(
                        str(row.get("countermeasure") or "-"),
                        styles["BodyText"],
                    ),
                    Paragraph(
                        str(row.get("responsible") or "-"),
                        styles["BodyText"],
                    ),
                    Paragraph(str(row.get("due_date") or "-"), styles["BodyText"]),
                    Paragraph(str(row.get("status") or "-"), styles["BodyText"]),
                    Paragraph(
                        str(row.get("standard_id") or "-"),
                        styles["BodyText"],
                    ),
                ]
            )

        follow_up_table_pdf = Table(
            table_rows,
            colWidths=[4.2 * cm, 4.7 * cm, 2.4 * cm, 1.9 * cm, 1.7 * cm, 2.1 * cm],
            repeatRows=1,
        )
        follow_up_table_pdf.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efe4ce")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#16302b")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9c9ab")),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#fffaf0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(follow_up_table_pdf)
        document.build(story)

    def export_follow_up_db_summary_jpg(output_path: Path) -> None:
        rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
        ]
        if not rows:
            raise ValueError(
                "Belum ada data follow up countermeasure untuk dibuatkan JPG."
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)

        def load_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
            font_candidates = []
            if bold:
                font_candidates.extend(
                    [
                        "arialbd.ttf",
                        "segoeuib.ttf",
                        str(Path("C:/Windows/Fonts/arialbd.ttf")),
                        str(Path("C:/Windows/Fonts/segoeuib.ttf")),
                    ]
                )
            else:
                font_candidates.extend(
                    [
                        "arial.ttf",
                        "segoeui.ttf",
                        str(Path("C:/Windows/Fonts/arial.ttf")),
                        str(Path("C:/Windows/Fonts/segoeui.ttf")),
                    ]
                )
            for font_name in font_candidates:
                try:
                    return ImageFont.truetype(font_name, size=size)
                except OSError:
                    continue
            return ImageFont.load_default()

        title_font = load_font(34, bold=True)
        section_title_font = load_font(22, bold=True)
        card_value_font = load_font(28, bold=True)
        card_label_font = load_font(14)
        body_font = load_font(16)
        small_font = load_font(13)
        mini_font = load_font(12)
        sample_draw = ImageDraw.Draw(Image.new("RGB", (10, 10), "white"))

        def line_height_for(font: ImageFont.ImageFont) -> int:
            bbox = sample_draw.textbbox((0, 0), "Ag", font=font)
            return bbox[3] - bbox[1]

        def wrap_text_to_width(
            text: object, max_width: int, font: ImageFont.ImageFont
        ) -> list[str]:
            normalized_text = str(text or "-").strip() or "-"
            wrapped_lines: list[str] = []
            for paragraph in normalized_text.splitlines() or [normalized_text]:
                words = paragraph.split()
                if not words:
                    wrapped_lines.append("")
                    continue
                current_line = words[0]
                for word in words[1:]:
                    candidate = f"{current_line} {word}"
                    if sample_draw.textlength(candidate, font=font) <= max_width:
                        current_line = candidate
                    else:
                        wrapped_lines.append(current_line)
                        current_line = word
                wrapped_lines.append(current_line)
            return wrapped_lines or ["-"]

        def parse_due_date(value: str) -> datetime | None:
            normalized = str(value or "").strip()
            if not normalized or normalized == "-":
                return None
            for date_format in (
                "%Y-%m-%d",
                "%d/%m/%Y",
                "%d-%m-%Y",
                "%m/%d/%Y",
                "%d %b %Y",
                "%d %B %Y",
            ):
                try:
                    return datetime.strptime(normalized, date_format)
                except ValueError:
                    continue
            return None

        def normalize_status(value: str) -> str:
            normalized = str(value or "").strip().casefold()
            if normalized in {"close", "closed", "complete", "done"}:
                return "Close"
            if normalized in {"open", "ongoing", "progress", "in progress"}:
                return "Open"
            return "Other"

        def draw_card(
            draw: ImageDraw.ImageDraw,
            left: int,
            top: int,
            width: int,
            height: int,
            label: str,
            value: str,
            accent: str,
            tint: str,
        ) -> None:
            draw.rounded_rectangle(
                (left, top, left + width, top + height),
                radius=24,
                fill="#fffaf0",
                outline="#ddcfb6",
                width=2,
            )
            draw.rounded_rectangle(
                (left + 18, top + 18, left + 34, top + height - 18),
                radius=12,
                fill=accent,
            )
            draw.rounded_rectangle(
                (left + 52, top + 18, left + width - 18, top + 54),
                radius=16,
                fill=tint,
            )
            draw.text(
                (left + 64, top + 27), label, fill="#6b4f3a", font=card_label_font
            )
            draw.text(
                (left + 52, top + 70), value, fill="#16302b", font=card_value_font
            )

        def draw_bar(
            draw: ImageDraw.ImageDraw,
            left: int,
            top: int,
            width: int,
            height: int,
            label: str,
            value: int,
            total: int,
            fill: str,
        ) -> None:
            draw.text((left, top), label, fill="#16302b", font=body_font)
            bar_top = top + 26
            draw.rounded_rectangle(
                (left, bar_top, left + width, bar_top + height),
                radius=height // 2,
                fill="#eadfcb",
            )
            bar_width = 0 if total <= 0 else max(24, int((value / total) * width))
            if value <= 0:
                bar_width = 0
            if bar_width:
                draw.rounded_rectangle(
                    (left, bar_top, left + min(bar_width, width), bar_top + height),
                    radius=height // 2,
                    fill=fill,
                )
            value_text = f"{value}"
            value_width = sample_draw.textlength(value_text, font=body_font)
            draw.text(
                (left + width - value_width, top),
                value_text,
                fill="#6b4f3a",
                font=body_font,
            )

        def draw_section_card(
            draw: ImageDraw.ImageDraw,
            left: int,
            top: int,
            right: int,
            bottom: int,
        ) -> None:
            draw.rounded_rectangle(
                (left, top, right, bottom),
                radius=26,
                fill="#fffaf0",
                outline="#ddcfb6",
                width=2,
            )

        today = datetime.now().date()
        soon_limit = today + timedelta(days=7)
        normalized_rows: list[dict[str, object]] = []
        for row in rows:
            status_category = normalize_status(str(row.get("status") or ""))
            due_date = parse_due_date(str(row.get("due_date") or ""))
            normalized_rows.append(
                {
                    **row,
                    "status_category": status_category,
                    "parsed_due_date": due_date,
                }
            )

        total_count = len(normalized_rows)
        close_count = sum(
            1 for row in normalized_rows if row["status_category"] == "Close"
        )
        open_count = sum(
            1 for row in normalized_rows if row["status_category"] == "Open"
        )
        overdue_count = sum(
            1
            for row in normalized_rows
            if row["status_category"] != "Close"
            and row["parsed_due_date"] is not None
            and row["parsed_due_date"].date() < today
        )
        due_soon_count = sum(
            1
            for row in normalized_rows
            if row["status_category"] != "Close"
            and row["parsed_due_date"] is not None
            and today <= row["parsed_due_date"].date() <= soon_limit
        )
        no_due_count = sum(
            1 for row in normalized_rows if row["parsed_due_date"] is None
        )

        responsible_counter = Counter(
            str(row.get("responsible") or "-")
            for row in normalized_rows
            if str(row.get("responsible") or "-").strip()
            and str(row.get("responsible") or "-") != "-"
        )
        file_counter = Counter(
            str(row.get("nama_file") or "-") for row in normalized_rows
        )
        top_responsibles = responsible_counter.most_common(5)
        top_files = file_counter.most_common(4)

        priority_rows = sorted(
            normalized_rows,
            key=lambda row: (
                row["status_category"] == "Close",
                row["parsed_due_date"] is None,
                row["parsed_due_date"] or datetime.max,
                str(row.get("nama_file") or ""),
            ),
        )[:5]

        image_width = 1600
        image_height = 1580 + max(0, len(priority_rows) - 3) * 124
        image = Image.new("RGB", (image_width, image_height), "#efe6d2")
        draw = ImageDraw.Draw(image)

        draw.ellipse((-120, -90, 460, 360), fill="#ead7b3")
        draw.ellipse((1120, -80, 1720, 360), fill="#d9e7da")
        draw.rounded_rectangle(
            (24, 24, image_width - 24, image_height - 24),
            radius=36,
            fill="#f7efe0",
        )

        draw.rounded_rectangle((52, 48, 1548, 190), radius=30, fill="#16302b")
        draw.rounded_rectangle((78, 74, 160, 156), radius=24, fill="#efe4ce")
        draw.text((97, 101), "IPS", fill="#16302b", font=section_title_font)
        draw.text(
            (188, 80),
            "Follow Up Countermeasure",
            fill="#fffaf0",
            font=title_font,
        )
        draw.text(
            (188, 126),
            f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            fill="#eadfcb",
            font=body_font,
        )

        card_top = 224
        card_width = 338
        card_gap = 24
        metrics_left = 88
        draw_card(
            draw,
            metrics_left,
            card_top,
            card_width,
            136,
            "Total Follow Up",
            str(total_count),
            "#0f5c4d",
            "#e3efe8",
        )
        draw_card(
            draw,
            metrics_left + card_width + card_gap,
            card_top,
            card_width,
            136,
            "Open Items",
            str(open_count),
            "#c96f3b",
            "#f5e6d7",
        )
        draw_card(
            draw,
            metrics_left + (card_width + card_gap) * 2,
            card_top,
            card_width,
            136,
            "Closed Items",
            str(close_count),
            "#245a4a",
            "#e1efe7",
        )
        draw_card(
            draw,
            metrics_left + (card_width + card_gap) * 3,
            card_top,
            card_width,
            136,
            "Overdue",
            str(overdue_count),
            "#a63d40",
            "#f6dfdf",
        )

        status_card = (74, 392, 756, 660)
        deadline_card = (74, 684, 756, 876)
        priority_card = (74, 924, 756, image_height - 72)
        ownership_card = (814, 392, 1526, 876)
        hotspot_card = (814, 924, 1526, image_height - 72)

        draw_section_card(draw, *status_card)
        draw_section_card(draw, *deadline_card)
        draw_section_card(draw, *priority_card)
        draw_section_card(draw, *ownership_card)
        draw_section_card(draw, *hotspot_card)

        draw.text(
            (104, 426), "Status Snapshot", fill="#16302b", font=section_title_font
        )
        draw.text(
            (104, 458),
            "Distribusi data aktif dari tabel follow up countermeasure",
            fill="#6b4f3a",
            font=small_font,
        )

        bar_left = 104
        bar_top = 510
        bar_width = 600
        status_total = max(total_count, 1)
        draw_bar(
            draw,
            bar_left,
            bar_top,
            bar_width,
            20,
            "Open",
            open_count,
            status_total,
            "#c96f3b",
        )
        draw_bar(
            draw,
            bar_left,
            bar_top + 76,
            bar_width,
            20,
            "Close",
            close_count,
            status_total,
            "#245a4a",
        )

        draw.text(
            (104, 718), "Deadline Outlook", fill="#16302b", font=section_title_font
        )
        deadline_cards = [
            ("Due Soon (7d)", due_soon_count, "#c96f3b", "#f5e6d7"),
            ("No Due Date", no_due_count, "#6b4f3a", "#efe4ce"),
            ("Need Attention", overdue_count + due_soon_count, "#a63d40", "#f6dfdf"),
        ]
        deadline_left = 104
        for index, (label, value, accent, tint) in enumerate(deadline_cards):
            left = deadline_left + index * 194
            draw.rounded_rectangle(
                (left, 760, left + 176, 856),
                radius=22,
                fill=tint,
                outline="#ddcfb6",
            )
            draw.rounded_rectangle(
                (left + 16, 776, left + 46, 806), radius=12, fill=accent
            )
            draw.text(
                (left + 16, 812), str(value), fill="#16302b", font=card_value_font
            )
            label_lines = wrap_text_to_width(label, 118, body_font)
            line_y = 772
            for line in label_lines:
                draw.text((left + 58, line_y), line, fill="#6b4f3a", font=body_font)
                line_y += line_height_for(body_font) + 2

        draw.text((844, 426), "Ownership Load", fill="#16302b", font=section_title_font)
        draw.text(
            (844, 458),
            "Responsible dengan jumlah tindak lanjut terbanyak",
            fill="#6b4f3a",
            font=small_font,
        )
        if top_responsibles:
            max_responsible = max(count for _, count in top_responsibles)
            row_top = 514
            for name, count in top_responsibles:
                label = name if name.strip() else "-"
                draw_bar(
                    draw,
                    844,
                    row_top,
                    620,
                    18,
                    label,
                    count,
                    max_responsible,
                    "#0f5c4d",
                )
                row_top += 74
        else:
            draw.text(
                (844, 522),
                "Belum ada responsible yang tercatat.",
                fill="#6b4f3a",
                font=body_font,
            )

        hotspot_top = 956
        draw.text(
            (844, hotspot_top), "File Hotspot", fill="#16302b", font=section_title_font
        )
        draw.text(
            (844, hotspot_top + 32),
            "File dengan jumlah countermeasure terbanyak",
            fill="#6b4f3a",
            font=small_font,
        )
        if top_files:
            max_files = max(count for _, count in top_files)
            row_top = hotspot_top + 88
            for file_name, count in top_files:
                draw_bar(
                    draw, 844, row_top, 620, 18, file_name, count, max_files, "#c96f3b"
                )
                row_top += 74
        else:
            draw.text(
                (844, hotspot_top + 88),
                "Belum ada data file follow up.",
                fill="#6b4f3a",
                font=body_font,
            )

        priority_panel_top = 948
        if priority_rows:
            draw.text(
                (104, priority_panel_top),
                "Priority Follow Up",
                fill="#16302b",
                font=section_title_font,
            )
            draw.text(
                (104, priority_panel_top + 32),
                "Item terbuka dengan due date paling dekat atau sudah lewat.",
                fill="#6b4f3a",
                font=small_font,
            )
            item_top = priority_panel_top + 92
            for index, row in enumerate(priority_rows, start=1):
                item_bottom = item_top + 98
                badge_color = (
                    "#245a4a" if row["status_category"] == "Close" else "#c96f3b"
                )
                if (
                    row["parsed_due_date"] is not None
                    and row["parsed_due_date"].date() < today
                ):
                    badge_color = "#a63d40"
                draw.rounded_rectangle(
                    (104, item_top, 726, item_bottom),
                    radius=18,
                    fill="#fcf5e8" if index % 2 == 0 else "#fffaf0",
                    outline="#eadfcb",
                )
                draw.rounded_rectangle(
                    (122, item_top + 31, 162, item_top + 67),
                    radius=10,
                    fill=badge_color,
                )
                draw.text(
                    (136, item_top + 38), str(index), fill="#fffaf0", font=small_font
                )
                title_lines = wrap_text_to_width(
                    row.get("countermeasure"), 446, body_font
                )[:2]
                title_y = item_top + 16
                for line in title_lines:
                    draw.text((180, title_y), line, fill="#16302b", font=body_font)
                    title_y += line_height_for(body_font) + 4
                meta_text = (
                    f"{row.get('nama_file') or '-'} | {row.get('responsible') or '-'} | "
                    f"Due {row.get('due_date') or '-'} | {row.get('status') or '-'}"
                )
                meta_lines = wrap_text_to_width(meta_text, 490, mini_font)[:2]
                meta_y = item_top + 64
                for line in meta_lines:
                    draw.text((180, meta_y), line, fill="#6b4f3a", font=mini_font)
                    meta_y += line_height_for(mini_font) + 3
                item_top += 114

        image.save(output_path, format="JPEG", quality=95)

    def export_follow_up_db_summary_excel(output_path: Path) -> None:
        rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
        ]
        if not rows:
            raise ValueError(
                "Belum ada data follow up countermeasure untuk dibuatkan Excel."
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Follow Up"
        worksheet.append(
            [
                "Nama File",
                "Countermeasure",
                "Responsible",
                "Due Date",
                "Status",
                "Standard ID",
            ]
        )
        for row in rows:
            worksheet.append(
                [
                    str(row.get("nama_file") or "-"),
                    str(row.get("countermeasure") or "-"),
                    str(row.get("responsible") or "-"),
                    str(row.get("due_date") or "-"),
                    str(row.get("status") or "-"),
                    str(row.get("standard_id") or "-"),
                ]
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        column_widths = {
            "A": 32,
            "B": 52,
            "C": 22,
            "D": 16,
            "E": 14,
            "F": 18,
        }
        for column_letter, width in column_widths.items():
            worksheet.column_dimensions[column_letter].width = width

        workbook.save(output_path)

    def has_exportable_data(notify_if_empty: bool = True) -> bool:
        current_view = str(app_state["result_view"])
        if current_view == "follow-up-db":
            has_rows = bool(app_state["follow_up_db_rows"])
            if not has_rows and notify_if_empty:
                notify(
                    "Belum ada data follow up countermeasure untuk diexport.",
                    error=True,
                )
            return has_rows

        has_results = bool(table_state["results"])
        if not has_results and notify_if_empty:
            notify("Belum ada hasil pemeriksaan untuk diexport.", error=True)
        return has_results

    def export_results_summary_pdf(output_path: Path) -> None:
        results = list(table_state["results"])
        if not results:
            raise ValueError("Belum ada hasil pemeriksaan untuk dibuatkan PDF.")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=1.5 * cm,
            rightMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
        )
        styles = getSampleStyleSheet()
        story = [
            Paragraph("IPS Completeness Summary", styles["Title"]),
            Spacer(1, 0.3 * cm),
            Paragraph(
                summary_text.value or "Belum ada ringkasan hasil.", styles["Normal"]
            ),
            Spacer(1, 0.4 * cm),
        ]

        table_rows = [["File Name", "Participant", "Status", "Completeness %"]]
        for result in results:
            status_label = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
            table_rows.append(
                [
                    Paragraph(str(result["file_name"]), styles["BodyText"]),
                    Paragraph(
                        str(result.get("participant") or "-"), styles["BodyText"]
                    ),
                    Paragraph(status_label, styles["BodyText"]),
                    Paragraph(
                        f"{float(result['completeness_percentage']):.2f}%",
                        styles["BodyText"],
                    ),
                ]
            )

        results_table_pdf = Table(
            table_rows, colWidths=[8.5 * cm, 3.5 * cm, 3.0 * cm, 3.0 * cm], repeatRows=1
        )
        results_table_pdf.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efe4ce")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#16302b")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9c9ab")),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#fffaf0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(results_table_pdf)
        document.build(story)

    def export_results_summary_jpg(output_path: Path) -> None:
        results = list(table_state["results"])
        if not results:
            raise ValueError("Belum ada hasil pemeriksaan untuk dibuatkan JPG.")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        base_font = ImageFont.load_default()
        title_font = ImageFont.load_default()
        image_width = 1600
        margin_x = 56
        top_margin = 44
        card_gap = 18
        row_gap = 6
        line_gap = 6

        sample_draw = ImageDraw.Draw(Image.new("RGB", (10, 10), "white"))

        def line_height_for(font: ImageFont.ImageFont) -> int:
            bbox = sample_draw.textbbox((0, 0), "Ag", font=font)
            return bbox[3] - bbox[1]

        def wrap_text_to_width(
            text: object, max_width: int, font: ImageFont.ImageFont
        ) -> list[str]:
            normalized_text = str(text or "-").strip() or "-"
            wrapped_lines: list[str] = []
            for paragraph in normalized_text.splitlines() or [normalized_text]:
                words = paragraph.split()
                if not words:
                    wrapped_lines.append("")
                    continue
                current_line = words[0]
                for word in words[1:]:
                    candidate = f"{current_line} {word}"
                    if sample_draw.textlength(candidate, font=font) <= max_width:
                        current_line = candidate
                    else:
                        wrapped_lines.append(current_line)
                        current_line = word
                wrapped_lines.append(current_line)
            return wrapped_lines or ["-"]

        def draw_text_lines(
            draw: ImageDraw.ImageDraw,
            lines: list[str],
            x: int,
            y: int,
            font: ImageFont.ImageFont,
            fill: str,
            extra_gap: int = line_gap,
        ) -> int:
            current_y = y
            text_line_height = line_height_for(font)
            for line in lines:
                draw.text((x, current_y), line, fill=fill, font=font)
                current_y += text_line_height + extra_gap
            return current_y

        def draw_card(
            draw: ImageDraw.ImageDraw,
            left: int,
            top: int,
            width: int,
            height: int,
            title: str,
            value: str,
            accent_color: str,
        ) -> None:
            draw.rounded_rectangle(
                (left, top, left + width, top + height),
                radius=24,
                fill="#fffaf0",
                outline="#d9c9ab",
                width=2,
            )
            draw.rounded_rectangle(
                (left + 18, top + 18, left + 58, top + 58), radius=14, fill=accent_color
            )
            draw.text((left + 78, top + 18), title, fill="#6b4f3a", font=base_font)
            draw.text((left + 78, top + 52), value, fill="#16302b", font=title_font)

        complete_count = sum(1 for result in results if result["is_complete"])
        incomplete_count = len(results) - complete_count
        average_completeness = sum(
            float(result["completeness_percentage"]) for result in results
        ) / len(results)

        summary_value = summary_text.value or (
            f"{len(results)} file diperiksa | {incomplete_count} file incomplete | {complete_count} file complete"
        )
        summary_lines = wrap_text_to_width(
            summary_value, image_width - (margin_x * 2) - 40, base_font
        )

        columns = [
            ("No", 70),
            ("File Name", 620),
            ("Participant", 270),
            ("Status", 220),
            ("Completeness", 220),
        ]
        table_left = margin_x
        table_width = sum(width for _, width in columns)
        header_height = 170
        cards_top = top_margin + header_height - 18
        card_width = (table_width - (card_gap * 3)) // 4
        card_height = 108
        table_top = cards_top + card_height + 32
        body_line_height = line_height_for(base_font)

        prepared_rows: list[dict[str, object]] = []
        for index, result in enumerate(results, start=1):
            status_label = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
            completeness_label = f"{float(result['completeness_percentage']):.2f}%"
            row_cells = [
                wrap_text_to_width(str(index), columns[0][1] - 24, base_font),
                wrap_text_to_width(result["file_name"], columns[1][1] - 24, base_font),
                wrap_text_to_width(
                    result.get("participant") or "-", columns[2][1] - 24, base_font
                ),
                wrap_text_to_width(status_label, columns[3][1] - 24, base_font),
                wrap_text_to_width(completeness_label, columns[4][1] - 24, base_font),
            ]
            max_lines = max(len(cell_lines) for cell_lines in row_cells)
            row_height = max(44, 18 + max_lines * (body_line_height + line_gap))
            prepared_rows.append(
                {
                    "cells": row_cells,
                    "status": status_label,
                    "row_height": row_height,
                }
            )

        table_header_height = 48
        footer_height = 48
        image_height = (
            table_top
            + table_header_height
            + sum(int(row["row_height"]) + row_gap for row in prepared_rows)
            + footer_height
            + 30
        )

        image = Image.new("RGB", (image_width, max(900, image_height)), "#f4f1e8")
        draw = ImageDraw.Draw(image)

        draw.rounded_rectangle(
            (24, 24, image_width - 24, image_height - 24), radius=32, fill="#f7efe0"
        )
        draw.rounded_rectangle(
            (margin_x, top_margin, image_width - margin_x, top_margin + header_height),
            radius=32,
            fill="#efe4ce",
        )
        draw.rounded_rectangle(
            (margin_x + 24, top_margin + 24, margin_x + 92, top_margin + 92),
            radius=20,
            fill="#0f5c4d",
        )
        draw.ellipse(
            (margin_x + 70, top_margin + 62, margin_x + 94, top_margin + 86),
            fill="#c96f3b",
        )
        draw.text(
            (margin_x + 118, top_margin + 26),
            "IPS Completeness Summary",
            fill="#16302b",
            font=title_font,
        )
        draw.text(
            (margin_x + 118, top_margin + 58),
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            fill="#6b4f3a",
            font=base_font,
        )
        draw_text_lines(
            draw,
            summary_lines,
            margin_x + 118,
            top_margin + 88,
            base_font,
            "#4b5d58",
            extra_gap=4,
        )

        draw_card(
            draw,
            table_left,
            cards_top,
            card_width,
            card_height,
            "Total File",
            str(len(results)),
            "#0f5c4d",
        )
        draw_card(
            draw,
            table_left + card_width + card_gap,
            cards_top,
            card_width,
            card_height,
            "Complete",
            str(complete_count),
            "#245a4a",
        )
        draw_card(
            draw,
            table_left + (card_width + card_gap) * 2,
            cards_top,
            card_width,
            card_height,
            "Incomplete",
            str(incomplete_count),
            "#a63d40",
        )
        draw_card(
            draw,
            table_left + (card_width + card_gap) * 3,
            cards_top,
            card_width,
            card_height,
            "Avg Completeness",
            f"{average_completeness:.1f}%",
            "#c96f3b",
        )

        draw.rounded_rectangle(
            (
                table_left,
                table_top,
                table_left + table_width,
                table_top + table_header_height,
            ),
            radius=20,
            fill="#16302b",
        )
        current_x = table_left
        for column_name, column_width in columns:
            draw.text(
                (current_x + 14, table_top + 15),
                column_name,
                fill="#fffaf0",
                font=base_font,
            )
            current_x += column_width

        current_y = table_top + table_header_height + 10
        for row_index, row in enumerate(prepared_rows):
            row_height = int(row["row_height"])
            status_label = str(row["status"])
            status_bg = "#e3f1ea" if status_label == "COMPLETE" else "#f7dfdf"
            status_fg = "#245a4a" if status_label == "COMPLETE" else "#a63d40"
            row_bg = "#fffaf0" if row_index % 2 == 0 else "#fcf5e8"
            draw.rounded_rectangle(
                (
                    table_left,
                    current_y,
                    table_left + table_width,
                    current_y + row_height,
                ),
                radius=18,
                fill=row_bg,
                outline="#e0d2b7",
            )

            current_x = table_left
            for cell_index, ((_, column_width), cell_lines) in enumerate(
                zip(columns, row["cells"])
            ):
                cell_x = current_x + 14
                cell_y = current_y + 12
                if cell_index == 3:
                    badge_width = min(
                        column_width - 28,
                        max(
                            110,
                            int(sample_draw.textlength(status_label, font=base_font))
                            + 38,
                        ),
                    )
                    badge_height = 30
                    draw.rounded_rectangle(
                        (cell_x, cell_y, cell_x + badge_width, cell_y + badge_height),
                        radius=15,
                        fill=status_bg,
                    )
                    draw.text(
                        (cell_x + 14, cell_y + 8),
                        status_label,
                        fill=status_fg,
                        font=base_font,
                    )
                else:
                    draw_text_lines(
                        draw,
                        list(cell_lines),
                        cell_x,
                        cell_y,
                        base_font,
                        "#16302b",
                        extra_gap=4,
                    )
                current_x += column_width

            current_y += row_height + row_gap

        draw.text(
            (table_left, current_y + 14),
            "Generated by IPS Checker",
            fill="#9a8874",
            font=base_font,
        )
        image.save(output_path, format="JPEG", quality=95)

    def request_export_path(
        dialog_title: str, extension: str, initial_filename: str | None = None
    ) -> Path | None:
        set_picker_open(True)
        page.update()
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        default_name = (
            initial_filename or f"ips-completeness-summary-{timestamp}.{extension}"
        )
        selected_path = with_hidden_tk_root(
            lambda _: filedialog.asksaveasfilename(
                title=dialog_title,
                initialdir=resolve_initial_directory(
                    str(app_state["csv_output"]), APP_DIR
                ),
                initialfile=default_name,
                defaultextension=f".{extension}",
                filetypes=[(f"{extension.upper()} files", f"*.{extension}")],
            )
        )
        set_picker_open(False)
        page.update()
        if not selected_path:
            return None
        app_state["csv_output"] = selected_path
        persist_config()
        return Path(selected_path)

    def open_workbook_file_action(result: dict[str, object]) -> None:
        workbook_path = str(result.get("workbook_path") or "")
        if not workbook_path:
            notify("Path file workbook tidak ditemukan.", error=True)
            return

        workbook_file = Path(workbook_path)
        if not workbook_file.exists():
            notify(f"File tidak ditemukan: {workbook_file.name}", error=True)
            return

        try:
            os.startfile(str(workbook_file))
        except Exception as exc:
            log_path = write_error_log("Buka file workbook gagal", exc)
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        notify(f"Membuka file {result['file_name']}.")

    def open_ips_generator(_: ft.ControlEvent) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return

        if not IPS_HTML_FILE.exists():
            notify(f"File tidak ditemukan: {IPS_HTML_FILE.name}", error=True)
            return

        try:
            os.startfile(str(IPS_HTML_FILE))
        except Exception as exc:
            log_path = write_error_log("Buka file IPS generator gagal", exc)
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        notify("Membuka IPS Generator.")

    def export_results_by_type(export_type: str) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return
        if not table_state["results"]:
            notify("Belum ada hasil pemeriksaan untuk diexport.", error=True)
            return

        export_type = export_type.lower()
        export_map = {
            "pdf": ("Simpan summary PDF", "pdf", export_results_summary_pdf),
            "jpg": ("Simpan summary JPG", "jpg", export_results_summary_jpg),
        }
        if export_type not in export_map:
            notify(f"Tipe export tidak dikenali: {export_type}", error=True)
            return

        dialog_title, extension, exporter = export_map[export_type]
        output_path = request_export_path(dialog_title, extension)
        if output_path is None:
            return

        try:
            exporter(output_path)
        except Exception as exc:
            log_path = write_error_log(
                f"Export {export_type.upper()} summary gagal", exc
            )
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        notify(f"Summary {export_type.upper()} berhasil dibuat.")

    def export_follow_up_db_by_type(export_type: str) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return
        if not app_state["follow_up_db_rows"]:
            notify(
                "Belum ada data follow up countermeasure untuk diexport.",
                error=True,
            )
            return

        export_type = export_type.lower()
        export_map = {
            "pdf": (
                "Simpan follow up PDF",
                "pdf",
                export_follow_up_db_summary_pdf,
            ),
            "jpg": (
                "Simpan follow up JPG",
                "jpg",
                export_follow_up_db_summary_jpg,
            ),
            "excel": (
                "Simpan follow up Excel",
                "xlsx",
                export_follow_up_db_summary_excel,
            ),
        }
        if export_type not in export_map:
            notify(f"Tipe export tidak dikenali: {export_type}", error=True)
            return

        dialog_title, extension, exporter = export_map[export_type]
        output_path = request_export_path(dialog_title, extension)
        if output_path is None:
            return

        try:
            exporter(output_path)
        except Exception as exc:
            log_path = write_error_log(
                f"Export {export_type.upper()} follow up gagal", exc
            )
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        notify(f"Follow up {export_type.upper()} berhasil dibuat.")

    def export_active_view_by_type(export_type: str) -> None:
        if str(app_state["result_view"]) == "follow-up-db":
            export_follow_up_db_by_type(export_type)
            return
        export_results_by_type(export_type)

    def handle_export_option(export_type: str) -> None:
        if getattr(export_dialog, "open", False):
            export_dialog.open = False
            page.pop_dialog()
            page.update()
        export_active_view_by_type(export_type)

    def handle_export_qr_option() -> None:
        if getattr(export_dialog, "open", False):
            export_dialog.open = False
            page.pop_dialog()
            page.update()
        if str(app_state["result_view"]) == "follow-up-db":
            show_follow_up_db_share_dialog()
            return
        show_results_share_dialog(None)

    def handle_follow_up_excel_option() -> None:
        if getattr(export_dialog, "open", False):
            export_dialog.open = False
            page.pop_dialog()
            page.update()
        export_follow_up_db_by_type("excel")

    def show_export_dialog(_: ft.ControlEvent) -> None:
        if app_state["is_picker_open"] or app_state["is_loading"]:
            return
        if not has_exportable_data():
            return

        current_view = str(app_state["result_view"])
        if current_view == "follow-up-db":
            export_description = (
                "Pilih format file export untuk tabel follow up countermeasure."
            )
            export_dialog.title = ft.Text(
                "Export Follow Up Countermeasure",
                weight=ft.FontWeight.W_700,
                color="#16302b",
            )
            third_export_button = ft.Button(
                content=ft.Text("Excel", weight=ft.FontWeight.W_600),
                icon=ft.Icons.TABLE_VIEW_ROUNDED,
                width=130,
                style=ft.ButtonStyle(bgcolor="#6b4f3a", color="#fffaf0"),
                on_click=lambda _: handle_follow_up_excel_option(),
            )
        else:
            export_description = (
                "Pilih format file export untuk ringkasan hasil pemeriksaan."
            )
            export_dialog.title = ft.Text(
                "Export Hasil Pemeriksaan",
                weight=ft.FontWeight.W_700,
                color="#16302b",
            )
            third_export_button = ft.Button(
                content=ft.Text("QR Code", weight=ft.FontWeight.W_600),
                icon=ft.Icons.QR_CODE_2_ROUNDED,
                width=130,
                style=ft.ButtonStyle(bgcolor="#6b4f3a", color="#fffaf0"),
                on_click=lambda _: handle_export_qr_option(),
            )

        export_dialog.content = ft.Container(
            width=460,
            content=ft.Column(
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        export_description,
                        color="#4b5d58",
                    ),
                    ft.Row(
                        alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                        controls=[
                            ft.Button(
                                content=ft.Text("PDF", weight=ft.FontWeight.W_600),
                                icon=ft.Icons.PICTURE_AS_PDF_ROUNDED,
                                width=130,
                                style=ft.ButtonStyle(
                                    bgcolor="#0f5c4d", color="#fffaf0"
                                ),
                                on_click=lambda _: handle_export_option("pdf"),
                            ),
                            ft.Button(
                                content=ft.Text("JPG", weight=ft.FontWeight.W_600),
                                icon=ft.Icons.IMAGE_ROUNDED,
                                width=130,
                                style=ft.ButtonStyle(
                                    bgcolor="#c96f3b", color="#fffaf0"
                                ),
                                on_click=lambda _: handle_export_option("jpg"),
                            ),
                            third_export_button,
                        ],
                    ),
                ],
            ),
        )
        export_dialog.actions = [ft.TextButton("Tutup", on_click=close_export_dialog)]
        if getattr(export_dialog, "open", False):
            page.update()
            return
        export_dialog.open = True
        page.show_dialog(export_dialog)

    def build_qr_image_bytes(text: str) -> bytes:
        qr = qrcode.QRCode(border=2, box_size=8)
        qr.add_data(text)
        qr.make(fit=True)
        image = qr.make_image(fill_color="black", back_color="white")
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()

    def show_share_dialog(result: dict[str, object]) -> None:
        summary_text_value = build_detail_summary_text(result)
        qr_image_bytes = build_qr_image_bytes(summary_text_value)

        share_dialog.title = ft.Text(
            "Share Detail Completeness", weight=ft.FontWeight.W_700, color="#16302b"
        )

        share_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        result["file_name"],
                        weight=ft.FontWeight.W_700,
                        color="#16302b",
                        text_align=ft.TextAlign.CENTER,
                    ),
                    ft.Image(
                        src=qr_image_bytes, width=260, height=260, fit=ft.BoxFit.CONTAIN
                    ),
                    ft.Container(
                        bgcolor="#f7efe0",
                        border=ft.Border.all(1, "#d9c9ab"),
                        border_radius=12,
                        padding=12,
                        content=ft.Text(
                            summary_text_value,
                            color="#4b5d58",
                            selectable=True,
                            size=12,
                        ),
                    ),
                ],
            ),
        )
        share_dialog.actions = [
            ft.TextButton(
                "Copy Summary",
                icon=ft.Icons.CONTENT_COPY_ROUNDED,
                on_click=lambda _: page.run_task(
                    copy_summary_to_clipboard, summary_text_value
                ),
            ),
            ft.TextButton("Tutup", on_click=close_share_dialog),
        ]

        if getattr(detail_dialog, "open", False):
            detail_dialog.open = False
            page.pop_dialog()

        if getattr(share_dialog, "open", False):
            page.update()
            return

        share_dialog.open = True
        page.show_dialog(share_dialog)

    def show_follow_up_share_dialog(result: dict[str, object]) -> None:
        summary_text_value = build_follow_up_summary_text(result)
        qr_image_bytes = build_qr_image_bytes(summary_text_value)

        share_dialog.title = ft.Text(
            "Share Follow Up Section 1.5", weight=ft.FontWeight.W_700, color="#16302b"
        )
        share_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        result["file_name"],
                        weight=ft.FontWeight.W_700,
                        color="#16302b",
                        text_align=ft.TextAlign.CENTER,
                    ),
                    ft.Image(
                        src=qr_image_bytes, width=260, height=260, fit=ft.BoxFit.CONTAIN
                    ),
                    ft.Container(
                        bgcolor="#f7efe0",
                        border=ft.Border.all(1, "#d9c9ab"),
                        border_radius=12,
                        padding=12,
                        content=ft.Text(
                            summary_text_value,
                            color="#4b5d58",
                            selectable=True,
                            size=12,
                        ),
                    ),
                ],
            ),
        )
        share_dialog.actions = [
            ft.TextButton(
                "Copy Summary",
                icon=ft.Icons.CONTENT_COPY_ROUNDED,
                on_click=lambda _: page.run_task(
                    copy_summary_to_clipboard, summary_text_value
                ),
            ),
            ft.TextButton("Tutup", on_click=close_share_dialog),
        ]

        if getattr(follow_up_dialog, "open", False):
            follow_up_dialog.open = False
            page.pop_dialog()

        if getattr(share_dialog, "open", False):
            page.update()
            return

        share_dialog.open = True
        page.show_dialog(share_dialog)

    def show_results_share_dialog(_: ft.ControlEvent) -> None:
        incomplete_only = False
        if incomplete_only:
            filtered_results = [
                result
                for result in table_state["results"]
                if not result.get("is_complete")
            ]
            empty_message = "Belum ada file incomplete untuk dibagikan."
            dialog_title = "Share Daftar File INCOMPLETE"
            heading_text = "Daftar File INCOMPLETE"
        else:
            filtered_results = list(table_state["results"])
            empty_message = "Belum ada hasil pemeriksaan untuk dibagikan."
            dialog_title = "Share Daftar Data IPS"
            heading_text = "Daftar Data IPS"

        if not filtered_results:
            notify(empty_message, error=True)
            return

        summary_text_value = build_results_summary_text(incomplete_only)
        qr_image_bytes = build_qr_image_bytes(summary_text_value)

        share_dialog.title = ft.Text(
            dialog_title, weight=ft.FontWeight.W_700, color="#16302b"
        )
        share_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        heading_text,
                        weight=ft.FontWeight.W_700,
                        color="#16302b",
                        text_align=ft.TextAlign.CENTER,
                    ),
                    ft.Image(
                        src=qr_image_bytes, width=260, height=260, fit=ft.BoxFit.CONTAIN
                    ),
                    ft.Container(
                        bgcolor="#f7efe0",
                        border=ft.Border.all(1, "#d9c9ab"),
                        border_radius=12,
                        padding=12,
                        content=ft.Text(
                            summary_text_value,
                            color="#4b5d58",
                            selectable=True,
                            size=12,
                        ),
                    ),
                ],
            ),
        )
        share_dialog.actions = [
            ft.TextButton(
                "Copy Summary",
                icon=ft.Icons.CONTENT_COPY_ROUNDED,
                on_click=lambda _: page.run_task(
                    copy_summary_to_clipboard, summary_text_value
                ),
            ),
            ft.TextButton("Tutup", on_click=close_share_dialog),
        ]

        if getattr(detail_dialog, "open", False):
            detail_dialog.open = False
            page.pop_dialog()
        if getattr(follow_up_dialog, "open", False):
            follow_up_dialog.open = False
            page.pop_dialog()

        if getattr(share_dialog, "open", False):
            page.update()
            return

        share_dialog.open = True
        page.show_dialog(share_dialog)

    def show_follow_up_db_share_dialog(_: ft.ControlEvent | None = None) -> None:
        rows = [
            normalize_follow_up_database_row(row)
            for row in app_state["follow_up_db_rows"]
        ]
        if not rows:
            notify(
                "Belum ada data follow up countermeasure untuk dibagikan.",
                error=True,
            )
            return

        summary_text_value = build_follow_up_db_summary_text()
        qr_image_bytes = build_qr_image_bytes(summary_text_value)

        share_dialog.title = ft.Text(
            "Share Follow Up Countermeasure",
            weight=ft.FontWeight.W_700,
            color="#16302b",
        )
        share_dialog.content = ft.Container(
            width=520,
            content=ft.Column(
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        "Daftar Follow Up Countermeasure",
                        weight=ft.FontWeight.W_700,
                        color="#16302b",
                        text_align=ft.TextAlign.CENTER,
                    ),
                    ft.Image(
                        src=qr_image_bytes,
                        width=260,
                        height=260,
                        fit=ft.BoxFit.CONTAIN,
                    ),
                    ft.Container(
                        bgcolor="#f7efe0",
                        border=ft.Border.all(1, "#d9c9ab"),
                        border_radius=12,
                        padding=12,
                        content=ft.Text(
                            summary_text_value,
                            color="#4b5d58",
                            selectable=True,
                            size=12,
                        ),
                    ),
                ],
            ),
        )
        share_dialog.actions = [
            ft.TextButton(
                "Copy Summary",
                icon=ft.Icons.CONTENT_COPY_ROUNDED,
                on_click=lambda _: page.run_task(
                    copy_summary_to_clipboard, summary_text_value
                ),
            ),
            ft.TextButton("Tutup", on_click=close_share_dialog),
        ]

        if getattr(detail_dialog, "open", False):
            detail_dialog.open = False
            page.pop_dialog()
        if getattr(follow_up_dialog, "open", False):
            follow_up_dialog.open = False
            page.pop_dialog()

        if getattr(share_dialog, "open", False):
            page.update()
            return

        share_dialog.open = True
        page.show_dialog(share_dialog)

    def show_info_dialog(_: ft.ControlEvent) -> None:
        def info_bullet(icon, title: str, description: str) -> ft.Control:
            return ft.Row(
                spacing=10,
                vertical_alignment=ft.CrossAxisAlignment.START,
                controls=[
                    ft.Icon(icon, size=18, color="#0f5c4d"),
                    ft.Column(
                        spacing=2,
                        tight=True,
                        expand=True,
                        controls=[
                            ft.Text(
                                title,
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=13,
                            ),
                            ft.Text(description, color="#4b5d58", size=12),
                        ],
                    ),
                ],
            )

        info_dialog.content = ft.Container(
            width=620,
            content=ft.Column(
                spacing=12,
                tight=True,
                controls=[
                    ft.Text(
                        "Informasi GUI", weight=ft.FontWeight.W_700, color="#16302b"
                    ),
                    ft.Text(
                        "Aplikasi ini dipakai untuk memeriksa completeness form IPS dari satu file Excel atau satu folder file Excel.",
                        color="#4b5d58",
                    ),
                    ft.Divider(color="#d9c9ab"),
                    ft.Text(
                        "Cara menggunakan", weight=ft.FontWeight.W_700, color="#16302b"
                    ),
                    info_bullet(
                        ft.Icons.FOLDER_OPEN_ROUNDED,
                        "Pilih sumber file",
                        "Gunakan Pilih Folder untuk banyak file atau Pilih File untuk satu dokumen Excel.",
                    ),
                    info_bullet(
                        ft.Icons.DESCRIPTION_ROUNDED,
                        "Buat IPS",
                        "Gunakan tombol Buat IPS di panel kiri untuk membuka halaman ips.html dan menyiapkan dokumen IPS baru.",
                    ),
                    info_bullet(
                        ft.Icons.PLAY_ARROW_ROUNDED,
                        "Jalankan pemeriksaan",
                        "Setelah file atau folder dipilih, proses check berjalan otomatis dan hasil muncul di panel kanan.",
                    ),
                    info_bullet(
                        ft.Icons.VISIBILITY_ROUNDED,
                        "Tinjau hasil",
                        "Di tabel hasil pemeriksaan, gunakan icon mata untuk melihat detail completeness dan icon open untuk membuka file IPS Excel asli.",
                    ),
                    info_bullet(
                        ft.Icons.TABLE_CHART_ROUNDED,
                        "Pindah view",
                        "Gunakan tombol menu di AppBar panel kanan untuk berpindah antara tabel hasil pemeriksaan dan database follow up countermeasure lokal otomatis.",
                    ),
                    info_bullet(
                        ft.Icons.EDIT_NOTE_ROUNDED,
                        "Kelola follow up",
                        "Pada view follow up countermeasure, Anda bisa melihat detail, mengubah status Open atau Close, mengisi Standard ID saat Close, dan menghapus data follow up.",
                    ),
                    info_bullet(
                        ft.Icons.DOWNLOAD_ROUNDED,
                        "Export",
                        "Gunakan tombol Export di AppBar panel kanan. Export mengikuti view aktif: hasil pemeriksaan mendukung PDF, JPG, dan QR Code, sedangkan follow up countermeasure mendukung PDF, JPG infografis, dan Excel.",
                    ),
                    ft.Divider(color="#d9c9ab"),
                    ft.Text(
                        "Dokumen IPS dianggap complete jika",
                        weight=ft.FontWeight.W_700,
                        color="#16302b",
                    ),
                    info_bullet(
                        ft.Icons.CHECK_CIRCLE_OUTLINE_ROUNDED,
                        "Section 1.1 dan 1.2 valid",
                        "Semua field wajib terisi, dengan trigger 1.1 cukup salah satu dari C6, C8, C10, atau C12.",
                    ),
                    info_bullet(
                        ft.Icons.RULE_FOLDER_OUTLINED,
                        "Section 1.3 valid",
                        "Setiap item hanya boleh punya satu nilai OK, NOK, atau NA, dan cell lain harus kosong.",
                    ),
                    info_bullet(
                        ft.Icons.TABLE_ROWS_ROUNDED,
                        "Section 1.4 valid",
                        "Minimal satu row penuh pada 66, 68, 70, 72, 74, 76, atau 78. Row terisi lain tidak boleh parsial.",
                    ),
                    info_bullet(
                        ft.Icons.TABLE_ROWS_ROUNDED,
                        "Section 1.5 valid",
                        "Minimal satu row penuh pada 88, 97, 106, 115, 124, atau 133. Row terisi lain tidak boleh parsial.",
                    ),
                    info_bullet(
                        ft.Icons.TASK_ALT_ROUNDED,
                        "Status COMPLETE",
                        "File dinyatakan complete jika semua rule section terpenuhi tanpa field atau row rule yang gagal.",
                    ),
                ],
            ),
        )
        info_dialog.actions = [ft.TextButton("Tutup", on_click=close_info_dialog)]
        if getattr(info_dialog, "open", False):
            page.update()
            return
        info_dialog.open = True
        page.show_dialog(info_dialog)

    def show_result_details(result: dict[str, object]) -> None:
        status_label = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
        section_completeness = result.get("section_completeness", {})

        section_table_rows = [
            ft.Container(
                bgcolor="#efe4ce",
                padding=ft.Padding.symmetric(horizontal=12, vertical=8),
                content=ft.Row(
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    controls=[
                        ft.Container(
                            expand=True,
                            content=ft.Text(
                                "Section",
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=12,
                            ),
                        ),
                        ft.Container(
                            width=90,
                            content=ft.Text(
                                "%",
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=12,
                            ),
                        ),
                    ],
                ),
            )
        ]
        for index, section in enumerate(sorted(section_completeness)):
            percentage = float(section_completeness.get(section, 0.0))
            section_table_rows.append(
                ft.Container(
                    bgcolor="#fffaf0" if index % 2 == 0 else "#fcf5e8",
                    border=ft.Border(bottom=ft.BorderSide(1, "#eadfcb")),
                    padding=ft.Padding.symmetric(horizontal=12, vertical=8),
                    content=ft.Row(
                        vertical_alignment=ft.CrossAxisAlignment.CENTER,
                        controls=[
                            ft.Container(
                                expand=True,
                                content=ft.Text(
                                    section_labels.get(str(section), str(section)),
                                    weight=ft.FontWeight.W_600,
                                    color="#16302b",
                                    size=12,
                                ),
                            ),
                            ft.Container(
                                width=90,
                                content=ft.Text(
                                    f"{percentage:.1f}%",
                                    color="#16302b",
                                    weight=ft.FontWeight.W_600,
                                    size=12,
                                ),
                            ),
                        ],
                    ),
                )
            )

        detail_controls: list[ft.Control] = [
            ft.Text(
                f"File: {result['file_name']}",
                color="#16302b",
                weight=ft.FontWeight.W_700,
            ),
            ft.Text(
                f"Participant: {result.get('participant') or '-'}", color="#4b5d58"
            ),
            ft.Text(f"Status: {status_label}", color="#4b5d58"),
            ft.Text(
                f"Completeness: {result['completeness_percentage']:.2f}%",
                color="#4b5d58",
            ),
            ft.Text("Per Section", weight=ft.FontWeight.W_700, color="#16302b"),
            ft.Container(
                border=ft.Border.all(1, "#d9c9ab"),
                border_radius=14,
                clip_behavior=ft.ClipBehavior.HARD_EDGE,
                content=ft.Column(spacing=0, controls=section_table_rows),
            ),
        ]

        detail_dialog.content = ft.Container(
            width=760,
            content=ft.Column(spacing=10, tight=True, controls=detail_controls),
        )
        detail_dialog.actions = [
            ft.TextButton(
                "Share",
                icon=ft.Icons.SHARE_ROUNDED,
                on_click=lambda _, item=result: show_share_dialog(item),
            ),
            ft.TextButton("Tutup", on_click=close_detail_dialog),
        ]
        if getattr(detail_dialog, "open", False):
            page.update()
            return
        detail_dialog.open = True
        page.show_dialog(detail_dialog)

    def show_follow_up_details(result: dict[str, object]) -> None:
        follow_up_rows = result.get("follow_up_rows", [])
        follow_up_table_rows = [
            ft.Container(
                bgcolor="#efe4ce",
                padding=ft.Padding.symmetric(horizontal=12, vertical=8),
                content=ft.Row(
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    controls=[
                        ft.Container(
                            expand=2,
                            content=ft.Text(
                                "Countermeasure",
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=12,
                            ),
                        ),
                        ft.Container(
                            expand=1,
                            content=ft.Text(
                                "Responsible",
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=12,
                            ),
                        ),
                        ft.Container(
                            width=120,
                            content=ft.Text(
                                "Due Date",
                                weight=ft.FontWeight.W_700,
                                color="#16302b",
                                size=12,
                            ),
                        ),
                    ],
                ),
            )
        ]
        if follow_up_rows:
            for index, row in enumerate(follow_up_rows):
                follow_up_table_rows.append(
                    ft.Container(
                        bgcolor="#fffaf0" if index % 2 == 0 else "#fcf5e8",
                        border=ft.Border(bottom=ft.BorderSide(1, "#eadfcb")),
                        padding=ft.Padding.symmetric(horizontal=12, vertical=8),
                        content=ft.Row(
                            vertical_alignment=ft.CrossAxisAlignment.START,
                            controls=[
                                ft.Container(
                                    expand=2,
                                    content=ft.Text(
                                        str(row.get("countermeasure") or "-"),
                                        color="#16302b",
                                        size=12,
                                    ),
                                ),
                                ft.Container(
                                    expand=1,
                                    content=ft.Text(
                                        str(row.get("responsible") or "-"),
                                        color="#4b5d58",
                                        size=12,
                                    ),
                                ),
                                ft.Container(
                                    width=120,
                                    content=ft.Text(
                                        str(row.get("due_date") or "-"),
                                        color="#4b5d58",
                                        size=12,
                                    ),
                                ),
                            ],
                        ),
                    )
                )
        else:
            follow_up_table_rows.append(
                ft.Container(
                    padding=16,
                    content=ft.Text(
                        "Belum ada data follow up pada section 1.5.", color="#6b4f3a"
                    ),
                )
            )

        follow_up_dialog.content = ft.Container(
            width=760,
            content=ft.Column(
                spacing=10,
                tight=True,
                controls=[
                    ft.Text(
                        f"File: {result['file_name']}",
                        color="#16302b",
                        weight=ft.FontWeight.W_700,
                    ),
                    ft.Text("Section 1.5 - Ask WHY WHY", color="#4b5d58"),
                    ft.Container(
                        border=ft.Border.all(1, "#d9c9ab"),
                        border_radius=14,
                        clip_behavior=ft.ClipBehavior.HARD_EDGE,
                        content=ft.Column(spacing=0, controls=follow_up_table_rows),
                    ),
                ],
            ),
        )
        follow_up_dialog.actions = [
            ft.TextButton(
                "Share",
                icon=ft.Icons.SHARE_ROUNDED,
                on_click=lambda _, item=result: show_follow_up_share_dialog(item),
            ),
            ft.TextButton("Tutup", on_click=close_follow_up_dialog),
        ]
        if getattr(follow_up_dialog, "open", False):
            page.update()
            return
        follow_up_dialog.open = True
        page.show_dialog(follow_up_dialog)

    def show_follow_up_details_from_database_row(row: dict[str, object]) -> None:
        file_name = str(row.get("nama_file") or "").strip()
        if not file_name:
            notify("Nama file follow up tidak ditemukan.", error=True)
            return

        show_follow_up_details(
            {
                "file_name": file_name,
                "follow_up_rows": [
                    {
                        "countermeasure": row.get("countermeasure"),
                        "responsible": row.get("responsible"),
                        "due_date": row.get("due_date"),
                    }
                ],
            }
        )

    def sort_results(column_index: int, ascending: bool) -> None:
        table_state["sort_column_index"] = column_index
        table_state["sort_ascending"] = ascending

        results = list(table_state["results"])
        if column_index == 0:
            results.sort(
                key=lambda item: str(item["file_name"]).casefold(),
                reverse=not ascending,
            )
        elif column_index == 1:
            results.sort(
                key=lambda item: float(item["completeness_percentage"]),
                reverse=not ascending,
            )

        table_state["results"] = results
        refresh_results_table()

    def sort_follow_up_rows(column_index: int, ascending: bool) -> None:
        table_state["follow_up_sort_column_index"] = column_index
        table_state["follow_up_sort_ascending"] = ascending

        def string_key(row: dict[str, object], field_name: str) -> str:
            return str(row.get(field_name) or "").strip().casefold()

        rows = list(app_state["follow_up_db_rows"])
        if column_index == 0:
            rows.sort(
                key=lambda row: string_key(row, "nama_file"),
                reverse=not ascending,
            )
        elif column_index == 1:
            rows.sort(
                key=lambda row: string_key(row, "countermeasure"),
                reverse=not ascending,
            )
        elif column_index == 2:
            rows.sort(
                key=lambda row: string_key(row, "responsible"),
                reverse=not ascending,
            )
        elif column_index == 3:
            rows.sort(
                key=lambda row: string_key(row, "due_date"),
                reverse=not ascending,
            )
        elif column_index == 4:
            rows.sort(
                key=lambda row: string_key(row, "status"),
                reverse=not ascending,
            )
        elif column_index == 5:
            rows.sort(
                key=lambda row: string_key(row, "standard_id"),
                reverse=not ascending,
            )

        app_state["follow_up_db_rows"] = rows
        refresh_follow_up_db_table()

    def toggle_sort(column_index: int) -> None:
        if table_state["sort_column_index"] == column_index:
            sort_results(column_index, not table_state["sort_ascending"])
            return
        sort_results(column_index, True)

    def toggle_follow_up_sort(column_index: int) -> None:
        if table_state["follow_up_sort_column_index"] == column_index:
            sort_follow_up_rows(
                column_index, not table_state["follow_up_sort_ascending"]
            )
            return
        sort_follow_up_rows(column_index, True)

    def build_sortable_header_cell(
        label: str,
        active_column_index: int,
        is_ascending: bool,
        column_index: int | None,
        toggle_handler,
    ) -> ft.Control:
        sort_suffix = ""
        if column_index is not None and active_column_index == column_index:
            sort_suffix = "  ^" if is_ascending else "  v"

        label_text = ft.Text(
            f"{label}{sort_suffix}",
            weight=ft.FontWeight.W_700,
            color="#16302b",
            text_align=ft.TextAlign.CENTER
            if column_index is None and label == "Action"
            else ft.TextAlign.LEFT,
        )
        if column_index is None:
            if label == "Action":
                return ft.Row(
                    alignment=ft.MainAxisAlignment.CENTER, controls=[label_text]
                )
            return label_text
        return ft.TextButton(
            content=label_text,
            style=ft.ButtonStyle(
                padding=0,
                color="#16302b",
                overlay_color="#00000000",
            ),
            on_click=lambda _: toggle_handler(column_index),
        )

    def build_header_cell(label: str, column_index: int | None = None) -> ft.Control:
        return build_sortable_header_cell(
            label,
            table_state["sort_column_index"],
            table_state["sort_ascending"],
            column_index,
            toggle_sort,
        )

    def build_result_row(result: dict[str, object], row_index: int) -> ft.Control:
        row_bgcolor = "#fffaf0" if row_index % 2 == 0 else "#fcf5e8"
        return ft.Container(
            bgcolor=row_bgcolor,
            border=ft.Border(bottom=ft.BorderSide(1, "#eadfcb")),
            padding=ft.Padding.symmetric(horizontal=16, vertical=14),
            content=ft.Row(
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                controls=[
                    ft.Container(
                        expand=True,
                        content=ft.Text(
                            str(result["file_name"]),
                            color="#16302b",
                            weight=ft.FontWeight.W_600,
                            overflow=ft.TextOverflow.ELLIPSIS,
                            max_lines=1,
                        ),
                    ),
                    ft.Container(
                        width=190,
                        content=build_completeness_cell(result),
                    ),
                    ft.Container(
                        width=96,
                        content=ft.Row(
                            spacing=0,
                            alignment=ft.MainAxisAlignment.CENTER,
                            vertical_alignment=ft.CrossAxisAlignment.CENTER,
                            controls=[
                                ft.IconButton(
                                    icon=ft.Icons.VISIBILITY_ROUNDED,
                                    icon_color="#0f5c4d",
                                    icon_size=18,
                                    tooltip="Lihat detail",
                                    on_click=lambda _, item=result: show_result_details(
                                        item
                                    ),
                                ),
                                ft.IconButton(
                                    icon=ft.Icons.OPEN_IN_NEW_ROUNDED,
                                    icon_color="#0f5c4d",
                                    icon_size=18,
                                    tooltip="Open file",
                                    on_click=lambda _, item=result: (
                                        open_workbook_file_action(item)
                                    ),
                                ),
                            ],
                        ),
                    ),
                ],
            ),
        )

    def refresh_results_table() -> None:
        header_row = ft.Container(
            bgcolor="#efe4ce",
            padding=ft.Padding.symmetric(horizontal=16, vertical=14),
            content=ft.Row(
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                controls=[
                    ft.Container(
                        expand=True, content=build_header_cell("File Name", 0)
                    ),
                    ft.Container(
                        width=190, content=build_header_cell("Completeness %", 1)
                    ),
                    ft.Container(width=96, content=build_header_cell("Action")),
                ],
            ),
        )
        result_rows = [
            build_result_row(result, index)
            for index, result in enumerate(table_state["results"])
        ]
        if not result_rows:
            result_rows = [
                ft.Container(
                    padding=24,
                    content=ft.Text("Belum ada hasil pemeriksaan.", color="#6b4f3a"),
                )
            ]
        results_table.controls = [header_row, *result_rows]

    refresh_results_table()
    refresh_follow_up_db_table()
    load_follow_up_database(update_page=False)
    refresh_result_panel_view()

    def start_run_check(selected_target: str | None = None) -> None:
        if selected_target:
            app_state["target_path"] = selected_target
        persist_config()
        set_loading(True)
        table_state["results"] = []
        refresh_results_table()
        summary_text.value = ""
        page.update()
        page.run_task(run_check, selected_target)

    async def run_check(selected_target: str | None = None) -> None:

        try:
            target_path = selected_target or app_state["target_path"]
            results = await asyncio.to_thread(
                run_checks,
                target_path=target_path,
                field_file=app_state["field_file"],
                sheet_name=app_state["sheet_name"],
                csv_output=None,
                base_dir=APP_DIR,
            )
        except Exception as exc:
            log_path = write_error_log("Pemeriksaan gagal", exc)
            set_loading(False)
            page.update()
            notify(str(exc), error=True)
            notify(f"Log error tersimpan di: {log_path.name}", error=True)
            return

        incomplete_count = sum(1 for result in results if not result["is_complete"])
        for result in results:
            result["section_summary"] = " | ".join(
                f"Section {section}: {percentage:.2f}% ({result['total_fields_by_section'].get(section, 0) - result['missing_fields_by_section'].get(section, 0)}/{result['total_fields_by_section'].get(section, 0)})"
                for section, percentage in sorted(
                    result.get("section_completeness", {}).items()
                )
            )
        summary_text.value = (
            f"{len(results)} file diperiksa | {incomplete_count} file incomplete | "
            f"{len(results) - incomplete_count} file complete"
        )
        table_state["results"] = list(results)
        sync_follow_up_database_from_results(list(results))
        sort_results(table_state["sort_column_index"], table_state["sort_ascending"])
        set_loading(False)
        page.update()
        notify("Pemeriksaan selesai.")

    form_panel = ft.Container(
        expand=True,
        bgcolor="#fffaf0",
        border_radius=24,
        padding=24,
        content=ft.Column(
            expand=True,
            spacing=18,
            controls=[
                app_logo,
                ft.Text(
                    "Pilih folder atau file Excel. Pemeriksaan akan langsung berjalan otomatis.",
                    color="#6b4f3a",
                ),
                folder_button,
                file_button,
                create_ips_button,
                info_button,
                ft.Container(expand=True),
                ft.Divider(color="#e6dac6", height=1),
                ft.Container(
                    width=side_button_width,
                    padding=ft.Padding.only(top=4),
                    content=ft.Text(
                        "© 2026 rardyant",
                        size=10,
                        color="#9a8874",
                        text_align=ft.TextAlign.CENTER,
                    ),
                ),
            ],
        ),
    )

    result_panel = ft.Container(
        expand=True,
        bgcolor="#e9dfc9",
        border_radius=24,
        padding=24,
        content=ft.Column(
            expand=True,
            spacing=16,
            controls=[
                ft.Container(
                    bgcolor="#fffaf0",
                    border=ft.Border.all(1, "#d9c9ab"),
                    border_radius=18,
                    clip_behavior=ft.ClipBehavior.HARD_EDGE,
                    content=result_view_menu,
                ),
                result_view_content,
            ],
        ),
    )

    root_layout = ft.Container(
        expand=True,
        content=ft.Row(
            expand=True,
            spacing=24,
            vertical_alignment=ft.CrossAxisAlignment.STRETCH,
            controls=[
                ft.Container(expand=2, content=form_panel),
                ft.Container(expand=10, content=result_panel),
            ],
        ),
    )

    def refresh_layout_height() -> None:
        page_padding = page.padding if isinstance(page.padding, (int, float)) else 0
        content_height = page.height
        if not content_height:
            content_height = max(
                1,
                int(page.window_height or DEFAULT_WINDOW_HEIGHT) - 48,
            )
        available_height = max(1, int(content_height) - int(page_padding * 2))
        form_panel.height = available_height
        result_panel.height = available_height
        root_layout.height = available_height

    folder_button.on_click = pick_target_folder
    file_button.on_click = pick_target_file
    create_ips_button.on_click = open_ips_generator
    export_action_button.on_click = show_export_dialog
    info_button.on_click = show_info_dialog
    page.on_resize = handle_page_resize
    page.on_close = handle_page_close
    sync_window_size_from_page()
    persist_config()

    page.add(root_layout)
    refresh_layout_height()
    page.update()
    page.services.append(clipboard_service)


if __name__ == "__main__":
    try:
        ft.run(main)
    except Exception as exc:
        log_path = write_error_log("Aplikasi gagal dijalankan", exc)
        raise RuntimeError(f"Aplikasi gagal dijalankan. Lihat log: {log_path}") from exc
