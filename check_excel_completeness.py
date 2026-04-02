from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook


FIELD_PATTERN = re.compile(r"^(?P<name>.+?)\s*\((?P<refs>[^)]+)\)\s*$")
SECTION_PATTERN = re.compile(r"^section\s+(?P<section>.+)$", re.IGNORECASE)
SECTION_13_ALLOWED_VALUES = {"OK", "NOK", "NA"}
SECTION_14_CHECK_ROWS = {66, 68, 70, 72, 74, 76, 78}
SECTION_15_CHECK_ROWS = {88, 97, 106, 115, 124, 133}


@dataclass(frozen=True)
class FieldSpec:
    section: str
    name: str
    cells: tuple[str, ...]


def is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False


def normalize_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().upper()
    return str(value).strip().upper()


def stringify_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def expand_reference_group(group: str) -> tuple[str, ...]:
    refs: list[str] = []
    for part in [item.strip().upper() for item in group.split(",") if item.strip()]:
        if ":" in part:
            start, end = part.split(":", 1)
            refs.extend(expand_range(start, end))
        else:
            refs.append(part)
    return tuple(refs)


def expand_range(start_ref: str, end_ref: str) -> Sequence[str]:
    from openpyxl.utils.cell import column_index_from_string, get_column_letter

    start_col, start_row = split_ref(start_ref)
    end_col, end_row = split_ref(end_ref)

    start_col_idx = column_index_from_string(start_col)
    end_col_idx = column_index_from_string(end_col)

    refs: list[str] = []
    for col_idx in range(start_col_idx, end_col_idx + 1):
        for row_idx in range(start_row, end_row + 1):
            refs.append(f"{get_column_letter(col_idx)}{row_idx}")
    return refs


def split_ref(cell_ref: str) -> tuple[str, int]:
    match = re.fullmatch(r"([A-Z]+)(\d+)", cell_ref.upper())
    if not match:
        raise ValueError(f"Cell reference is invalid: {cell_ref}")
    return match.group(1), int(match.group(2))


def parse_field_file(field_file: Path) -> list[FieldSpec]:
    fields: list[FieldSpec] = []
    seen: set[tuple[str, str, tuple[str, ...]]] = set()
    current_section = "unknown"

    for raw_line in field_file.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue

        section_match = SECTION_PATTERN.match(line)
        if section_match:
            current_section = section_match.group("section")
            continue

        field_match = FIELD_PATTERN.match(line)
        if not field_match:
            continue

        field = FieldSpec(
            section=current_section,
            name=field_match.group("name").strip(),
            cells=expand_reference_group(field_match.group("refs")),
        )

        field_key = (field.section, field.name.casefold(), field.cells)
        if field_key in seen:
            continue

        seen.add(field_key)
        fields.append(field)

    if not fields:
        raise ValueError(f"No field definitions found in {field_file}")

    return fields


def filter_fields_to_rows(fields: Sequence[FieldSpec], allowed_rows: set[int]) -> list[FieldSpec]:
    filtered_fields: list[FieldSpec] = []
    for field in fields:
        filtered_cells = tuple(cell for cell in field.cells if split_ref(cell)[1] in allowed_rows)
        if not filtered_cells:
            continue
        filtered_fields.append(
            FieldSpec(
                section=field.section,
                name=field.name,
                cells=filtered_cells,
            )
        )
    return filtered_fields


def find_excel_files(target_path: Path) -> list[Path]:
    if target_path.is_file():
        return [target_path]

    files = sorted(path for path in target_path.glob("*.xlsx") if not path.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No .xlsx files found in {target_path}")
    return files


def evaluate_field(worksheet, field: FieldSpec) -> tuple[bool, list[str], str | None]:
    if field.section == "1.1" and field.name.casefold() == "trigger":
        filled_refs = [cell for cell in field.cells if not is_blank(worksheet[cell].value)]
        if filled_refs:
            return True, [], None
        return False, list(field.cells), "At least one trigger cell must be filled."

    if field.section != "1.3":
        empty_refs = [cell for cell in field.cells if is_blank(worksheet[cell].value)]
        if empty_refs:
            return False, empty_refs, None
        return True, [], None

    normalized_values = {cell: normalize_cell_value(worksheet[cell].value) for cell in field.cells}
    valid_filled_refs = [cell for cell, value in normalized_values.items() if value in SECTION_13_ALLOWED_VALUES]
    nonblank_invalid_refs = [
        cell for cell, value in normalized_values.items() if value and value not in SECTION_13_ALLOWED_VALUES
    ]
    blank_refs = [cell for cell, value in normalized_values.items() if not value]

    is_complete = len(valid_filled_refs) == 1 and len(blank_refs) == len(field.cells) - 1 and not nonblank_invalid_refs
    if is_complete:
        return True, [], None

    issue_cells = sorted(set(valid_filled_refs + nonblank_invalid_refs + blank_refs))
    issue_reason = (
        "Exactly one cell must contain OK, NOK, or NA, and the remaining cells must be blank."
    )
    return False, issue_cells, issue_reason


def evaluate_row_based_section(
    worksheet,
    section_fields: Sequence[FieldSpec],
    section_label: str,
) -> tuple[bool, list[str], str | None]:
    row_to_cells: dict[int, list[str]] = defaultdict(list)
    all_refs: list[str] = []

    for field in section_fields:
        for cell in field.cells:
            row_number = split_ref(cell)[1]
            row_to_cells[row_number].append(cell)
            all_refs.append(cell)

    expected_cells_per_row = len(section_fields)
    complete_rows: list[int] = []
    partial_rows: list[int] = []
    partial_row_refs: list[str] = []

    for row_number, cells in sorted(row_to_cells.items()):
        filled_cells = [cell for cell in cells if not is_blank(worksheet[cell].value)]
        if len(cells) != expected_cells_per_row:
            partial_rows.append(row_number)
            partial_row_refs.extend(cells)
            continue
        if len(filled_cells) == expected_cells_per_row:
            complete_rows.append(row_number)
        elif filled_cells:
            partial_rows.append(row_number)
            partial_row_refs.extend(cells)

    if complete_rows and not partial_rows:
        return True, [], None

    if partial_rows:
        issue_refs = sorted(set(partial_row_refs))
        issue_reason = (
            f"Section {section_label} requires at least one fully filled row, and any other populated row must also be fully filled or completely blank."
        )
        return False, issue_refs, issue_reason

    issue_reason = (
        f"At least one row in section {section_label} must have all required columns filled."
    )
    return False, sorted(all_refs), issue_reason


def extract_row_based_section_entries(
    worksheet,
    section_fields: Sequence[FieldSpec],
) -> list[dict[str, str]]:
    field_cells_by_name = {field.name.casefold(): field.cells for field in section_fields}
    countermeasure_cells = field_cells_by_name.get("countermeasure", ())
    responsible_cells = field_cells_by_name.get("responsible", ())
    due_date_cells = field_cells_by_name.get("due date", ())

    row_count = max(len(countermeasure_cells), len(responsible_cells), len(due_date_cells))
    entries: list[dict[str, str]] = []

    for index in range(row_count):
        countermeasure = stringify_cell_value(worksheet[countermeasure_cells[index]].value) if index < len(countermeasure_cells) else ""
        responsible = stringify_cell_value(worksheet[responsible_cells[index]].value) if index < len(responsible_cells) else ""
        due_date = stringify_cell_value(worksheet[due_date_cells[index]].value) if index < len(due_date_cells) else ""

        if not any((countermeasure, responsible, due_date)):
            continue

        row_number = ""
        if index < len(countermeasure_cells):
            row_number = str(split_ref(countermeasure_cells[index])[1])
        elif index < len(responsible_cells):
            row_number = str(split_ref(responsible_cells[index])[1])
        elif index < len(due_date_cells):
            row_number = str(split_ref(due_date_cells[index])[1])

        entries.append(
            {
                "row": row_number,
                "countermeasure": countermeasure,
                "responsible": responsible,
                "due_date": due_date,
            }
        )

    return entries


def check_workbook(workbook_path: Path, fields: Iterable[FieldSpec], sheet_name: str | None) -> dict[str, object]:
    field_list = list(fields)
    section_14_fields = filter_fields_to_rows(
        [field for field in field_list if field.section == "1.4"],
        SECTION_14_CHECK_ROWS,
    )
    section_15_fields = filter_fields_to_rows(
        [field for field in field_list if field.section == "1.5"],
        SECTION_15_CHECK_ROWS,
    )
    regular_fields = [field for field in field_list if field.section not in {"1.4", "1.5"}]
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        missing_fields: list[dict[str, object]] = []
        total_fields_by_section: dict[str, int] = defaultdict(int)
        missing_by_section: dict[str, int] = defaultdict(int)
        total_cells = 0
        missing_cells = 0

        for field in regular_fields:
            total_fields_by_section[field.section] += 1
            is_complete, issue_refs, issue_reason = evaluate_field(worksheet, field)
            total_cells += len(field.cells)
            missing_cells += len(issue_refs)
            if not is_complete:
                missing_by_section[field.section] += 1
                missing_fields.append(
                    {
                        "section": field.section,
                        "field": field.name,
                        "expected_cells": ", ".join(field.cells),
                        "missing_cells": ", ".join(issue_refs),
                        "issue": issue_reason,
                        "filled": False,
                    }
                )

        if section_14_fields:
            total_fields_by_section["1.4"] += 1
            total_cells += sum(len(field.cells) for field in section_14_fields)
            is_complete, issue_refs, issue_reason = evaluate_row_based_section(worksheet, section_14_fields, "1.4")
            missing_cells += len(issue_refs)
            if not is_complete:
                missing_by_section["1.4"] += 1
                missing_fields.append(
                    {
                        "section": "1.4",
                        "field": "row completeness",
                        "expected_cells": ", ".join(field.name for field in section_14_fields),
                        "missing_cells": ", ".join(issue_refs),
                        "issue": issue_reason,
                        "filled": False,
                    }
                )

        if section_15_fields:
            total_fields_by_section["1.5"] += 1
            total_cells += sum(len(field.cells) for field in section_15_fields)
            is_complete, issue_refs, issue_reason = evaluate_row_based_section(worksheet, section_15_fields, "1.5")
            missing_cells += len(issue_refs)
            if not is_complete:
                missing_by_section["1.5"] += 1
                missing_fields.append(
                    {
                        "section": "1.5",
                        "field": "row completeness",
                        "expected_cells": ", ".join(field.name for field in section_15_fields),
                        "missing_cells": ", ".join(issue_refs),
                        "issue": issue_reason,
                        "filled": False,
                    }
                )

        section_completeness = {
            section: round(((total_count - missing_by_section.get(section, 0)) / total_count) * 100, 2)
            for section, total_count in sorted(total_fields_by_section.items())
        }
        completeness_percentage = round(
            sum(section_completeness.values()) / len(section_completeness),
            2,
        ) if section_completeness else 0.0
        participant_value = worksheet["O10"].value
        follow_up_rows = extract_row_based_section_entries(worksheet, section_15_fields) if section_15_fields else []

        return {
            "file_name": workbook_path.name,
            "sheet_name": worksheet.title,
            "participant": "" if participant_value is None else str(participant_value).strip(),
            "total_fields": sum(total_fields_by_section.values()),
            "total_fields_by_section": dict(total_fields_by_section),
            "missing_fields_count": len(missing_fields),
            "missing_fields_by_section": dict(missing_by_section),
            "section_completeness": section_completeness,
            "completeness_percentage": completeness_percentage,
            "total_cells": total_cells,
            "missing_cells": missing_cells,
            "is_complete": len(missing_fields) == 0,
            "missing_fields": missing_fields,
            "follow_up_rows": follow_up_rows,
        }
    finally:
        workbook.close()


def build_section_summary(result: dict[str, object], separator: str = ", ") -> str:
    section_counts = result.get("missing_fields_by_section", {})
    if not section_counts:
        return "-"

    return separator.join(
        f"Section {section}: {count}"
        for section, count in sorted(section_counts.items())
    )


def build_section_completeness_summary(result: dict[str, object], separator: str = ", ") -> str:
    section_percentages = result.get("section_completeness", {})
    if not section_percentages:
        return "-"

    missing_by_section = result.get("missing_fields_by_section", {})
    total_fields_by_section = result.get("total_fields_by_section", {})
    return separator.join(
        (
            f"Section {section}: {percentage:.2f}% "
            f"({total_fields_by_section.get(section, 0) - missing_by_section.get(section, 0)}/"
            f"{total_fields_by_section.get(section, 0)})"
        )
        for section, percentage in sorted(section_percentages.items())
    )


def resolve_input_path(path_value: str | Path, base_dir: Path | None = None) -> Path:
    path = Path(path_value)
    if path.is_absolute():
        return path.resolve()

    if base_dir is None:
        base_dir = Path.cwd()
    return (base_dir / path).resolve()


def run_checks(
    target_path: str | Path,
    field_file: str | Path,
    sheet_name: str | None = None,
    csv_output: str | Path | None = None,
    base_dir: Path | None = None,
) -> list[dict[str, object]]:
    resolved_base_dir = base_dir or Path.cwd()
    resolved_target_path = resolve_input_path(target_path, resolved_base_dir)
    resolved_field_file = resolve_input_path(field_file, resolved_base_dir)

    if not resolved_field_file.exists():
        raise FileNotFoundError(f"Field file not found: {resolved_field_file}")
    if not resolved_target_path.exists():
        raise FileNotFoundError(f"Target path not found: {resolved_target_path}")

    fields = parse_field_file(resolved_field_file)
    workbooks = find_excel_files(resolved_target_path)
    results = [check_workbook(workbook_path, fields, sheet_name) for workbook_path in workbooks]

    if csv_output:
        output_file = resolve_input_path(csv_output, resolved_base_dir)
        write_csv_report(results, output_file)

    return results


def print_report(results: Sequence[dict[str, object]]) -> None:
    for result in results:
        status = "COMPLETE" if result["is_complete"] else "INCOMPLETE"
        print(f"\n[{status}] {result['file_name']}")
        print(f"  Missing fields: {result['missing_fields_count']}")
        print(f"  Completeness: {result['completeness_percentage']:.2f}%")

        if result["section_completeness"]:
            print(f"  Per section: {build_section_completeness_summary(result)}")


def write_csv_report(results: Sequence[dict[str, object]], output_file: Path) -> None:
    with output_file.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "file_name",
                "participant",
                "status",
                "missing_fields_count",
                "completeness_percentage",
                "section_summary",
            ],
        )
        writer.writeheader()

        for result in results:
            writer.writerow(
                {
                    "file_name": result["file_name"],
                    "participant": result.get("participant", ""),
                    "status": "COMPLETE" if result["is_complete"] else "INCOMPLETE",
                    "missing_fields_count": result["missing_fields_count"],
                    "completeness_percentage": f"{result['completeness_percentage']:.2f}",
                    "section_summary": build_section_completeness_summary(result, separator=" | "),
                }
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Check Excel form completeness based on required cells listed in a text file."
    )
    parser.add_argument(
        "target",
        nargs="?",
        default="new create ips",
        help="Target .xlsx file or folder containing .xlsx files. Default: 'new create ips'.",
    )
    parser.add_argument(
        "--field-file",
        default="field_cell.txt",
        help="Path to the field definition file. Default: field_cell.txt",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Worksheet name to inspect. Default: first sheet in workbook.",
    )
    parser.add_argument(
        "--csv",
        default=None,
        help="Optional CSV output file path for the report.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    base_dir = Path(__file__).resolve().parent
    results = run_checks(
        target_path=args.target,
        field_file=args.field_file,
        sheet_name=args.sheet,
        csv_output=args.csv,
        base_dir=base_dir,
    )

    print_report(results)

    if args.csv:
        output_file = resolve_input_path(args.csv, base_dir)
        print(f"\nCSV report written to: {output_file}")

    incomplete_count = sum(1 for result in results if not result["is_complete"])
    print(f"\nSummary: {incomplete_count} incomplete file(s) out of {len(results)} checked.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())